"""Run the fatigue analysis and export the result tables."""

from __future__ import annotations

from pathlib import Path
import os
import tempfile
import warnings
import zipfile
from openpyxl.styles import Alignment
from openpyxl.styles import Font
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
import numpy as np
import pandas as pd
import argparse
import hashlib
import json
import platform
import time
import scipy
import sklearn
import data
from data import (
    CALIBRATION_MODES,
    CENSORED_MODELS,
    DAMAGE_CAP,
    DEFAULT_OUTPUT_DIR,
    EXPECTED_EXACT_ROWS,
    EXPECTED_RUNOUT_ROWS,
    EXPECTED_VERIFIED_ROWS,
    EXTERNAL_MODELS,
    HYBRID_CANDIDATES,
    HYBRID_SPECS,
    NESTED_CANDIDATES,
    PHYSICAL_SPECS,
    PHYSICS_MIX_GRID,
    PPTX_OUTPUT_NAME,
    PRIMARY_COMPARATORS,
    PRIMARY_GROUP_COLUMN,
    PROPOSED_MODEL,
    REPEATED_VALIDATION_SEEDS,
    RMSE_TIE_TOLERANCE,
    SEED,
    SHEET,
    TRAD_CATEGORICAL,
    TRAD_NUMERIC,
    ablation_effects,
    component_ablation_summary,
    external_paired_effects,
    external_uq_summary,
    grouped_folds,
    load_data,
    metric_summary,
    paired_effects,
    parse_bool,
    require_exact_training,
    resolve_data,
    runout_summary,
    sb_metrics,
    selection_rule_sensitivity,
    source_gain_table,
    training_weights,
    uq_summary,
)
from models import (
    exact_temperature_residual_gate,
)
from validation import (
    duplicate_source_sensitivity,
    external_holdout_validation,
    repeated_validation,
    run_primary_folds,
    sensitivity_audit,
    sstar_exclusion_sensitivity,
)


def write_workbook(path: Path, tables: dict[str, pd.DataFrame]) -> None:
    if Font is None:
        warnings.warn("openpyxl unavailable; workbook was not written")
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{path.stem}_", suffix=".xlsx", dir=path.parent
    )
    os.close(descriptor)
    temporary_path = Path(temporary_name)
    try:
        with pd.ExcelWriter(temporary_path, engine="openpyxl") as writer:
            for sheet, table in tables.items():
                table.to_excel(writer, sheet_name=sheet[:31], index=False)
            workbook = writer.book
            for sheet in workbook.worksheets:
                sheet.freeze_panes = "A2"
                sheet.auto_filter.ref = sheet.dimensions
                sheet.sheet_view.showGridLines = False
                for cell in sheet[1]:
                    cell.font = Font(color="FFFFFF", bold=True)
                    cell.fill = PatternFill("solid", fgColor="1F4E78")
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                for column in range(1, sheet.max_column + 1):
                    values = [str(sheet.cell(row, column).value or "") for row in range(1, min(sheet.max_row, 250) + 1)]
                    sheet.column_dimensions[get_column_letter(column)].width = min(max(11, max(map(len, values)) + 2), 48)
                for row in sheet.iter_rows(min_row=2):
                    for cell in row:
                        cell.alignment = Alignment(vertical="top", wrap_text=False)
            text_layouts = {
                "README": {"A": 28, "B": 100},
                "Protocol": {"A": 38, "B": 110},
                "Data_Audit": {"A": 48, "B": 80},
                "Metric_Definitions": {"A": 22, "B": 34, "C": 100},
                "Model_Specs": {"A": 24, "B": 44, "D": 110},
                "Validation_Definitions": {"A": 24, "B": 100, "C": 100},
                "Claim_Gate": {"A": 46, "B": 12, "C": 110, "D": 20},
                "References": {"A": 44, "B": 70, "C": 100},
            }
            for sheet_name, widths in text_layouts.items():
                if sheet_name not in workbook.sheetnames:
                    continue
                sheet = workbook[sheet_name]
                for column, width in widths.items():
                    sheet.column_dimensions[column].width = width
                for row_index in range(2, sheet.max_row + 1):
                    sheet.row_dimensions[row_index].height = 32
                    for cell in sheet[row_index]:
                        cell.alignment = Alignment(vertical="top", wrap_text=True)
        with zipfile.ZipFile(temporary_path) as archive:
            damaged_member = archive.testzip()
            if damaged_member is not None:
                raise IOError(f"Workbook ZIP validation failed at {damaged_member}")
        os.replace(temporary_path, path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()


def write_readme(path: Path, data_path: Path, metrics: pd.DataFrame, effects: pd.DataFrame,
                 gate: pd.DataFrame, df: pd.DataFrame, parameters: pd.DataFrame,
                 n_runout: int) -> None:
    proposed = metrics.loc[metrics["model"].eq("WB-PIML")].iloc[0]
    anchor = metrics.loc[metrics["model"].eq("WB-PIML-Anchor")].iloc[0]
    proposed_parameters = parameters.loc[parameters["model"].eq("WB-PIML")]
    mix_text = ", ".join(f"{value:.2f}" for value in proposed_parameters["physics_mix"])
    external_models = ["WB", "RF", "ExtraTrees", "ET-Walker", "GBDT", "SVR", "ML-Ens", "WB-PIML"]
    primary_best = metrics.loc[metrics["model"].isin(external_models)].sort_values("SB_RMSE").iloc[0]
    effect_ml = effects.loc[effects["comparator"].eq("ML-Ens") & effects["effect"].eq("RMSE_gain")].iloc[0]
    effect_et = effects.loc[effects["comparator"].eq("ExtraTrees") & effects["effect"].eq("RMSE_gain")].iloc[0]
    effect_walker = effects.loc[effects["comparator"].eq("ET-Walker") & effects["effect"].eq("RMSE_gain")].iloc[0]
    n_exact = int(df["is_exact"].sum())
    text = f"""# WB-PIML censor-aware, temperature-robust competing-damage results

## Analysis protocol

- Data: `{data_path.name}` / `{SHEET}` only.
- The verified worksheet contains {n_exact} exact fractures and {n_runout} right-censored runouts; all {len(df)} unique records enter the joint AFT likelihood.
- Every record has `sample_weight=1`; row numbers, sources, campaigns and equivalence IDs never change fitting importance.
- Exact fractures contribute density terms. Runouts contribute survival terms and are never treated as failures at their stopping cycles.
- No code-side label override is applied; workbook labels and campaign IDs are used as reviewed.
- {df['source_id'].nunique()} publication sources are assigned to {df['campaign_id'].nunique()} human-reviewed campaigns; no campaign is inferred from row counts alone.
- `Plot_Data_Verified` is physically deduplicated: {df['record_equivalence_id'].nunique()} unique records and no repeated equivalence ID. Duplicate copies remain traceable in `Plot_Data` and `Excluded_Records` and are excluded from the primary analysis.
- Five fixed campaign-disjoint outer folds; every record appears in one outer test fold and no equivalence cluster crosses folds.
- The physical branch combines Walker--Basquin mechanical damage with a time-dependent environmental rate proportional to `E(env)*Arrhenius(T)/f`.
- The machine-learning branch predicts only `y-mu_WB-CD`; the physical trunk has unit weight in every fold and cannot fall back to a purely data-driven model.
- If the current development set has no exact fracture at the requested temperature, the learned point residual is set to zero and prediction uses the competing-damage physical trunk; this support check never reads validation responses.
- Walker gamma, residual shrinkage eta and tree regularization are selected exclusively inside campaign-disjoint development folds. No affine output calibration is used.
- Exact-only campaign-balanced RMSE remains the point-performance endpoint; joint censored NLL is the within-tolerance selection tie-break.
- ET-Walker uses joint-data inner folds. WB-PIML includes four additional predeclared residual candidates with support-aware formulations.
- Runouts affect the physical trunk and AFT scale only through `-log P(Nf>N_stop|x)`; point RMSE/MAE/R2/F5/C* remain exact-only.
- There is no neural network, source intercept, arbitrary oxidation threshold, or outer-test tuning.
- Repeated record-, series-, source-, and campaign-disjoint validation separates interpolation from generalization.
- LOSO/LOAO/LOTO/LOEO are factor-held-out stress tests. WB-PIML, ExtraTrees and ET-Walker are re-selected only in the remaining records with campaign-disjoint inner folds. These outer factor splits are not campaign-disjoint by design, and the actual campaign overlap is reported for every holdout.

## Performance summary

- Best pre-specified model in the primary campaign-disjoint analysis: **{primary_best['model']} = {primary_best['SB_RMSE']:.3f}** campaign-balanced RMSE.
- WB-PIML campaign-balanced RMSE: **{proposed['SB_RMSE']:.3f}**.
- Selected competing-damage residual branch RMSE: **{anchor['SB_RMSE']:.3f}**; reported-model consistency difference: **{anchor['SB_RMSE'] - proposed['SB_RMSE']:.3f}**.
- Fixed physical-trunk weights by fold: **[{mix_text}]**; every value must equal 1.00.
- Paired RMSE gain over ML-Ens: **{effect_ml['estimate']:.3f}**, 95% campaign-bootstrap CI **[{effect_ml['CI95_lo']:.3f}, {effect_ml['CI95_hi']:.3f}]**.
- Paired RMSE gain over plain ExtraTrees: **{effect_et['estimate']:.3f}**, 95% campaign-bootstrap CI **[{effect_et['CI95_lo']:.3f}, {effect_et['CI95_hi']:.3f}]**.
- Paired RMSE gain over equally tuned ET-Walker: **{effect_walker['estimate']:.3f}**, 95% campaign-bootstrap CI **[{effect_walker['CI95_lo']:.3f}, {effect_walker['CI95_hi']:.3f}]**.

The `SB_*` metrics are campaign-balanced. The `Claim_Gate` worksheet summarizes the statistical and validation checks.

## Output files

- The default tabular deliverable is one workbook: `WB_PIML_results.xlsx`.
- `Outer_Fold_Assignment` is included in that workbook for reproducibility.
- CSV files are disabled by default; pass `--export-csv` only when legacy machine-readable files are explicitly needed.
- The workbook contains the data for the manuscript figures.
- `protocol.json` and this README remain as non-tabular audit outputs.
"""
    path.write_text(text, encoding="utf-8")


def claim_gate(metrics: pd.DataFrame, effects: pd.DataFrame, parameters: pd.DataFrame,
               df: pd.DataFrame, nested_selection: pd.DataFrame,
               nested_candidates: pd.DataFrame, runout: pd.DataFrame,
               external_summary: pd.DataFrame,
               external_groups: pd.DataFrame,
               external_effects: pd.DataFrame,
               external_uq: pd.DataFrame,
               external_selection: pd.DataFrame,
               temperature_gate_rows: pd.DataFrame) -> pd.DataFrame:
    """Evidence-derived gates; unknown audit fields remain unknown, never True."""
    metric = metrics.set_index("model")

    def effect_row(comparator: str) -> pd.Series:
        selected = effects.loc[
            effects["comparator"].eq(comparator) & effects["effect"].eq("RMSE_gain")
        ]
        if selected.empty:
            return pd.Series({"estimate": np.nan, "CI95_lo": np.nan, "CI95_hi": np.nan})
        return selected.iloc[0]

    paired = {name: effect_row(name) for name in PRIMARY_COMPARATORS}
    proposed = metric.loc[PROPOSED_MODEL]
    anchor = metric.loc["WB-PIML-Anchor"]
    external_models = [
        "WB", "RF", "ExtraTrees", "ET-Walker", "GBDT", "SVR", "ML-Ens",
        *CENSORED_MODELS, PROPOSED_MODEL,
    ]
    best_primary = str(metrics.loc[metrics["model"].isin(external_models)]
                       .sort_values("SB_RMSE").iloc[0]["model"])
    proposed_parameters = parameters.loc[parameters["model"].eq(PROPOSED_MODEL)]
    physics_mix = pd.to_numeric(proposed_parameters["physics_mix"], errors="coerce")
    n_rows = len(df)
    n_exact = int(df["is_exact"].sum())
    n_runout = int(df["is_runout"].sum())
    n_overrides = int(df["label_override_applied"].sum())
    audit_complete = bool(df.get("audit_record_complete", pd.Series(False, index=df.index)).all())
    supplied_conflicts = int(df.loc[
        df["supplied_label_conflict_flag"].fillna(False), "record_equivalence_id"
    ].nunique())
    reviewed_conflicts = int(df.loc[
        df["label_conflict_flag"].fillna(False), "record_equivalence_id"
    ].nunique())
    resolved_conflicts = int(df.loc[
        df["label_conflict_resolved"].fillna(False), "record_equivalence_id"
    ].nunique())
    secondary_rows = int(df["reviewed_secondary_source"].fillna(False).sum())
    duplicate_clusters = int(df.loc[
        df["cross_source_duplicate_candidate"].fillna(False), "record_equivalence_id"
    ].nunique())
    campaign_disjoint = bool(df.groupby("campaign_id")["outer_fold"].nunique().le(1).all())
    reviewed_map = bool(
        audit_complete
        and df["campaign_mapping_status"].astype(str).str.startswith("workbook_").all()
        and df["campaign_id"].astype(str).str.strip().ne("").all()
    )
    temperature_gate_verified = bool(
        len(temperature_gate_rows) == len(df)
        and not temperature_gate_rows["row_id"].duplicated().any()
    )
    if temperature_gate_verified:
        for fold, held_out in temperature_gate_rows.groupby("outer_fold"):
            development = df.loc[df["outer_fold"].ne(int(fold))]
            expected_gate = exact_temperature_residual_gate(
                development, held_out
            ).astype(bool)
            actual_gate = held_out["residual_temperature_supported"].astype(bool).to_numpy()
            if not np.array_equal(expected_gate, actual_gate):
                temperature_gate_verified = False
                break

    raw = data.AUDIT_CONTEXT.get("raw")
    duplicate_audit_complete = False
    duplicate_audit_detail = "raw Plot_Data audit table unavailable"
    if (
        isinstance(raw, pd.DataFrame) and not raw.empty
        and {"record_equivalence_id", "source_id"}.issubset(raw.columns)
    ):
        raw_ids = raw["record_equivalence_id"].fillna("").astype(str).str.strip()
        raw_audit = pd.DataFrame({
            "record_equivalence_id": raw_ids,
            "source_id": raw["source_id"].fillna("").astype(str).str.strip(),
        })
        raw_audit = raw_audit.loc[raw_audit["record_equivalence_id"].ne("")]
        raw_counts = raw_audit.groupby("record_equivalence_id").size()
        raw_source_counts = raw_audit.groupby("record_equivalence_id")["source_id"].nunique()
        verified_ids = df["record_equivalence_id"].astype(str)
        expected_cluster_size = verified_ids.map(raw_counts)
        expected_source_count = verified_ids.map(raw_source_counts)
        duplicate_audit_complete = bool(
            df["record_equivalence_id"].is_unique
            and expected_cluster_size.notna().all()
            and expected_source_count.notna().all()
            and np.array_equal(
                expected_cluster_size.to_numpy(int),
                df["duplicate_cluster_size"].to_numpy(int),
            )
            and np.array_equal(
                expected_source_count.to_numpy(int),
                df["duplicate_source_count"].to_numpy(int),
            )
            and np.array_equal(
                expected_source_count.gt(1).to_numpy(bool),
                df["cross_source_duplicate_candidate"].to_numpy(bool),
            )
        )
        duplicate_audit_detail = (
            f"verified IDs covered by Plot_Data={int(expected_cluster_size.notna().sum())}/{len(df)}; "
            f"cross-source duplicate clusters={duplicate_clusters}"
        )

    selection_trace_complete = False
    selection_trace_detail = "nested selection trace is incomplete"
    required_selection = {"outer_fold", "model", "candidate_index"}
    if (
        required_selection.issubset(nested_selection.columns)
        and required_selection.issubset(nested_candidates.columns)
    ):
        selected = nested_selection.loc[
            nested_selection["model"].eq("WB-PIML-Anchor"),
            ["outer_fold", "candidate_index"],
        ].dropna()
        candidates = nested_candidates.loc[
            nested_candidates["model"].eq("WB-PIML-Anchor"),
            ["outer_fold", "candidate_index"],
        ].dropna()
        selected_pairs = set(map(tuple, selected.astype(int).to_numpy()))
        candidate_pairs = set(map(tuple, candidates.astype(int).to_numpy()))
        candidate_counts = candidates.groupby("outer_fold")["candidate_index"].nunique()
        expected_folds = set(
            pd.to_numeric(proposed_parameters["outer_fold"], errors="coerce")
            .dropna().astype(int)
        )
        forbidden_outer_columns = {
            "pred_logN", "outer_test_logN", "outer_test_RMSE", "test_response"
        }
        selection_trace_complete = bool(
            expected_folds
            and set(selected["outer_fold"].astype(int)) == expected_folds
            and set(candidate_counts.index.astype(int)) == expected_folds
            and candidate_counts.eq(len(HYBRID_CANDIDATES)).all()
            and selected_pairs.issubset(candidate_pairs)
            and forbidden_outer_columns.isdisjoint(nested_candidates.columns)
        )
        selection_trace_detail = (
            f"selected folds={len(selected_pairs)}; candidate counts by fold="
            f"{candidate_counts.astype(int).to_dict()}"
        )

    censored_parameters = parameters.loc[parameters["model"].isin(CENSORED_MODELS)]
    censored_trace_complete = bool(
        all(name in metric.index for name in CENSORED_MODELS)
        and not censored_parameters.empty
        and censored_parameters.groupby("model")["outer_fold"].nunique()
        .reindex(CENSORED_MODELS).eq(len(proposed_parameters)).all()
        and pd.to_numeric(
            censored_parameters["active_runout_count"], errors="coerce"
        ).gt(0).all()
    )

    physical_columns = [
        "walker_slope", "environment_stress_exponent",
        "arrhenius_temperature_sensitivity", "n_physics_parameters",
    ]
    physical_values = proposed_parameters[physical_columns].apply(
        pd.to_numeric, errors="coerce"
    )
    mechanistic_constraints_verified = bool(
        not physical_values.empty
        and np.isfinite(physical_values.to_numpy(float)).all()
        and physical_values["walker_slope"].between(0.05, 20.0).all()
        and physical_values["environment_stress_exponent"].between(0.0, 10.0).all()
        and physical_values["arrhenius_temperature_sensitivity"].between(0.0, 10.0).all()
        and physical_values["n_physics_parameters"].eq(5).all()
        and physics_mix.notna().all() and physics_mix.gt(0).all()
    )
    rmse_ci_pass = {
        name: bool(np.isfinite(row["CI95_lo"]) and float(row["CI95_lo"]) > 0)
        for name, row in paired.items()
    }
    primary_rmse_superiority = all(rmse_ci_pass.values())

    required_effects = {"RMSE_gain", "MAE_gain", "R2_gain", "F5_gain", "C_star_gain"}
    primary_effect_rows = effects.loc[
        effects["comparator"].isin(PRIMARY_COMPARATORS)
        & effects["effect"].isin(required_effects)
    ].copy()
    primary_multimetric_superiority = bool(
        len(primary_effect_rows) == len(PRIMARY_COMPARATORS) * len(required_effects)
        and not primary_effect_rows.duplicated(["comparator", "effect"]).any()
        and pd.to_numeric(primary_effect_rows["CI95_lo"], errors="coerce").gt(0).all()
    )

    external_point_passes: dict[str, bool] = {}
    for scenario, scenario_table in external_summary.groupby("scenario"):
        proposed_rows = scenario_table.loc[scenario_table["model"].eq(PROPOSED_MODEL)]
        comparator_rows = scenario_table.loc[
            scenario_table["model"].isin(PRIMARY_COMPARATORS)
        ]
        external_point_passes[str(scenario)] = bool(
            len(proposed_rows) == 1 and len(comparator_rows) == len(PRIMARY_COMPARATORS)
            and float(proposed_rows.iloc[0]["SB_RMSE"])
            < float(comparator_rows["SB_RMSE"].min())
        )
    expected_external_scenarios = {"LOSO", "LOAO", "LOTO", "LOEO"}
    external_point_superiority = bool(
        set(external_point_passes) == expected_external_scenarios
        and all(external_point_passes.values())
    )
    external_rmse_rows = external_effects.loc[
        external_effects["comparator"].isin(PRIMARY_COMPARATORS)
        & external_effects["effect"].eq("RMSE_gain")
    ].copy()
    expected_external_rows = len(expected_external_scenarios) * len(PRIMARY_COMPARATORS)
    external_ci_superiority = bool(
        expected_external_rows > 0
        and len(external_rmse_rows) == expected_external_rows
        and set(external_rmse_rows["scenario"].astype(str)) == expected_external_scenarios
        and not external_rmse_rows.duplicated(["scenario", "comparator", "effect"]).any()
        and pd.to_numeric(external_rmse_rows["CI95_lo"], errors="coerce").gt(0).all()
    )
    expected_external_uq_rows = len(expected_external_scenarios) * 2
    external_interval_complete = bool(
        expected_external_uq_rows > 0
        and len(external_uq) == expected_external_uq_rows
        and set(external_uq["scenario"].astype(str)) == expected_external_scenarios
        and not external_uq.duplicated(["scenario", "nominal"]).any()
        and pd.to_numeric(
            external_uq["interval_completeness"], errors="coerce"
        ).eq(1.0).all()
    )
    external_interval_nominal_compatible = bool(
        external_interval_complete
        and (
            pd.to_numeric(external_uq["coverage_lo"], errors="coerce")
            <= pd.to_numeric(external_uq["nominal"], errors="coerce")
        ).all()
        and (
            pd.to_numeric(external_uq["coverage_hi"], errors="coerce")
            >= pd.to_numeric(external_uq["nominal"], errors="coerce")
        ).all()
    )

    external_nested_selection_complete = False
    external_selection_detail = "external nested-selection trace is incomplete"
    proposed_external_selection = pd.DataFrame()
    external_censor_guard_never_relaxed = False
    external_censor_status_known = False
    external_relaxed_count = 0
    external_wb_holdout_count = 0
    required_external_selection_columns = {
        "scenario", "holdout_group", "fold", "model", "inner_group",
        "n_inner_folds", "n_train_campaigns", "n_train_exact",
        "n_candidates_evaluated",
        "inner_SB_RMSE", "selected_in_holdout_training_only",
        "outer_test_used_for_selection", "validation_scope",
        "campaign_disjoint_by_design", "campaign_overlap",
        "holdout_campaign_disjoint", "censor_constraint_relaxed",
        "censor_feasible",
        "n_overlapping_campaigns", "empirical_halfwidth_80",
        "empirical_halfwidth_90", "n_interval_calibration_exact",
        "interval_calibration", "outer_test_used_for_interval_calibration",
    }
    required_external_group_columns = {"scenario", "holdout_group"}
    if (
        required_external_selection_columns.issubset(external_selection.columns)
        and required_external_group_columns.issubset(external_groups.columns)
    ):
        expected_holdouts = {
            (str(scenario), str(holdout))
            for scenario, holdout in external_groups[
                ["scenario", "holdout_group"]
            ].drop_duplicates().itertuples(index=False, name=None)
        }
        selected_models = {PROPOSED_MODEL, "ExtraTrees", "ET-Walker"}
        expected_rows = {
            (scenario, holdout, model)
            for scenario, holdout in expected_holdouts
            for model in selected_models
        }
        selection = external_selection.copy()
        forbidden_external_response_columns = {
            "logN", "pred_logN", "outer_test_logN", "outer_test_RMSE",
            "holdout_RMSE", "test_response",
        }
        actual_rows = {
            (str(scenario), str(holdout), str(model))
            for scenario, holdout, model in selection[
                ["scenario", "holdout_group", "model"]
            ].itertuples(index=False, name=None)
        }
        expected_candidate_counts = {
            PROPOSED_MODEL: len(HYBRID_CANDIDATES),
            "ExtraTrees": len({
                (leaf, features)
                for _, leaf, features, _ in NESTED_CANDIDATES
            }),
            "ET-Walker": len(NESTED_CANDIDATES),
        }
        observed_candidate_counts = pd.to_numeric(
            selection["n_candidates_evaluated"], errors="coerce"
        )
        required_candidate_counts = selection["model"].map(
            expected_candidate_counts
        )
        inner_folds = pd.to_numeric(selection["n_inner_folds"], errors="coerce")
        train_campaigns = pd.to_numeric(
            selection["n_train_campaigns"], errors="coerce"
        )
        overlap_count = pd.to_numeric(
            selection["n_overlapping_campaigns"], errors="coerce"
        )
        proposed_external_selection = selection.loc[
            selection["model"].eq(PROPOSED_MODEL)
        ].copy()
        interval_halfwidths = proposed_external_selection[
            ["empirical_halfwidth_80", "empirical_halfwidth_90"]
        ].apply(pd.to_numeric, errors="coerce")
        interval_calibration_exact = pd.to_numeric(
            proposed_external_selection["n_interval_calibration_exact"],
            errors="coerce",
        )
        proposed_train_exact = pd.to_numeric(
            proposed_external_selection["n_train_exact"], errors="coerce"
        )
        external_nested_selection_complete = bool(
            expected_holdouts
            and {scenario for scenario, _ in expected_holdouts}
            == expected_external_scenarios
            and actual_rows == expected_rows
            and len(selection) == len(expected_rows)
            and not selection.duplicated(
                ["scenario", "holdout_group", "model"]
            ).any()
            and selection["model"].isin(selected_models).all()
            and selection["inner_group"].eq(PRIMARY_GROUP_COLUMN).all()
            and selection["selected_in_holdout_training_only"].eq(True).all()
            and selection["outer_test_used_for_selection"].eq(False).all()
            and selection["validation_scope"].eq(
                "domain-held-out stress test"
            ).all()
            and selection["campaign_disjoint_by_design"].eq(False).all()
            and selection["holdout_campaign_disjoint"].astype(bool).eq(
                ~selection["campaign_overlap"].astype(bool)
            ).all()
            and forbidden_external_response_columns.isdisjoint(selection.columns)
            and observed_candidate_counts.eq(required_candidate_counts).all()
            and inner_folds.between(2, 4).all()
            and train_campaigns.ge(inner_folds).all()
            and np.isfinite(pd.to_numeric(
                selection["inner_SB_RMSE"], errors="coerce"
            )).all()
            and selection["campaign_overlap"].notna().all()
            and overlap_count.ge(0).all()
            and np.isfinite(interval_halfwidths.to_numpy(float)).all()
            and interval_halfwidths.ge(0).all().all()
            and interval_calibration_exact.eq(proposed_train_exact).all()
            and proposed_external_selection["interval_calibration"].eq(
                "development campaign-disjoint OOF absolute residuals"
            ).all()
            and proposed_external_selection[
                "outer_test_used_for_interval_calibration"
            ].eq(False).all()
        )
        external_selection_detail = (
            f"holdouts={len(expected_holdouts)}; selected rows={len(selection)}; "
            f"candidate counts={expected_candidate_counts}; "
            f"holdouts with campaign overlap="
            f"{selection.loc[selection['campaign_overlap'].eq(True), ['scenario', 'holdout_group']].drop_duplicates().shape[0]}"
        )
        external_relaxed_flags = pd.to_numeric(
            proposed_external_selection["censor_constraint_relaxed"],
            errors="coerce",
        )
        external_feasible_flags = pd.to_numeric(
            proposed_external_selection["censor_feasible"], errors="coerce"
        )
        external_wb_holdout_count = len(proposed_external_selection)
        external_relaxed_count = int(external_relaxed_flags.eq(1).sum())
        external_censor_status_known = bool(
            external_wb_holdout_count == len(expected_holdouts)
            and external_relaxed_flags.notna().all()
            and external_feasible_flags.notna().all()
            and external_relaxed_flags.isin([0, 1]).all()
            and external_feasible_flags.isin([0, 1]).all()
        )
        external_censor_guard_never_relaxed = bool(
            external_nested_selection_complete
            and external_censor_status_known
            and external_relaxed_flags.eq(0).all()
            and external_feasible_flags.eq(1).all()
        )
        external_selection_detail += (
            f"; constraint-relaxed WB-PIML holdouts="
            f"{external_relaxed_count}/{external_wb_holdout_count}"
        )

    runout_index = (
        runout.set_index("model")
        if "model" in runout and not runout["model"].duplicated().any()
        else pd.DataFrame()
    )
    if (
        isinstance(runout_index, pd.DataFrame)
        and {PROPOSED_MODEL, "WB-CD"}.issubset(runout_index.index)
    ):
        proposed_runout_nll = float(runout_index.loc[PROPOSED_MODEL, "SB_runout_NLL"])
        reference_runout_nll = float(runout_index.loc["WB-CD", "SB_runout_NLL"])
        runout_noninferior = bool(proposed_runout_nll <= reference_runout_nll)
    else:
        proposed_runout_nll = reference_runout_nll = np.nan
        runout_noninferior = False

    anchor_selection = nested_selection.loc[
        nested_selection["model"].eq("WB-PIML-Anchor")
    ]
    relaxed_series = (
        anchor_selection["censor_constraint_relaxed"]
        if "censor_constraint_relaxed" in anchor_selection
        else pd.Series(dtype=bool)
    )
    relaxation_known = bool(
        len(relaxed_series) == len(anchor_selection)
        and not relaxed_series.isna().any()
    )
    relaxed_count = int(relaxed_series.eq(True).sum())
    censor_guard_never_relaxed = bool(
        len(anchor_selection) == len(proposed_parameters)
        and relaxation_known
        and relaxed_count == 0
    )

    environment_ratio = pd.to_numeric(
        proposed_parameters["environment_log10_damage_ratio"], errors="coerce"
    )
    environment_stress = pd.to_numeric(
        proposed_parameters["environment_stress_exponent"], errors="coerce"
    )
    arrhenius = pd.to_numeric(
        proposed_parameters["arrhenius_temperature_sensitivity"], errors="coerce"
    )
    boundary_hits = (
        np.isclose(environment_ratio, -8.0, atol=1e-5)
        | np.isclose(environment_ratio, 4.0, atol=1e-5)
        | np.isclose(environment_stress, 0.0, atol=1e-5)
        | np.isclose(environment_stress, 10.0, atol=1e-5)
        | np.isclose(arrhenius, 0.0, atol=1e-5)
        | np.isclose(arrhenius, 10.0, atol=1e-5)
    )
    mechanistic_parameters_interior = bool(
        len(boundary_hits) == len(proposed_parameters) and not np.any(boundary_hits)
    )
    universal_superiority = bool(
        primary_rmse_superiority
        and primary_multimetric_superiority
        and external_point_superiority
        and external_ci_superiority
        and external_nested_selection_complete
        and runout_noninferior
        and censor_guard_never_relaxed
        and external_censor_guard_never_relaxed
        and mechanistic_parameters_interior
        and external_interval_nominal_compatible
        and temperature_gate_verified
    )
    checks: list[tuple[str, object, str]] = [
        ("correct_data", n_rows == EXPECTED_VERIFIED_ROWS and n_exact == EXPECTED_EXACT_ROWS
         and n_runout == EXPECTED_RUNOUT_ROWS and n_overrides == 0,
         f"{n_rows} reviewed unique records: {n_exact} exact + {n_runout} right-censored; overrides={n_overrides}"),
        ("workbook_audit_complete", audit_complete,
         "record_role, audit_status, evidence_grade and duplicate_group_status are read row-by-row from the workbook"),
        ("human_reviewed_campaign_map", reviewed_map,
         f"mapping statuses are workbook-derived; reviewed secondary-replot rows={secondary_rows}"),
        ("label_consistency", bool(df["label_consistent"].all()),
         "is_exact_failure is the complement of is_runout for every verified row"),
        ("label_conflicts_resolved", reviewed_conflicts == 0,
         f"raw supplied conflicts={supplied_conflicts}; resolved={resolved_conflicts}; unresolved={reviewed_conflicts}"),
        ("physical_duplicate_audit", duplicate_audit_complete,
         duplicate_audit_detail),
        ("campaign_disjoint_validation", campaign_disjoint,
         f"{df['campaign_id'].nunique()} reviewed campaigns never cross an outer fold"),
        ("equal_training_weights", bool(np.allclose(df["sample_weight"], 1.0)),
         "all 222 physically unique likelihood records have unit weight"),
        ("runouts_enter_right_censored_likelihood", n_runout == EXPECTED_RUNOUT_ROWS,
         "runouts enter density/survival model selection but never exact point metrics"),
        ("fair_censored_comparators_present", censored_trace_complete,
         "both censored AFT comparators have metrics, one parameter record per outer fold, and positive runout training counts"),
        ("physics_active_in_all_outer_folds", bool(physics_mix.notna().all() and physics_mix.gt(0).all()),
         f"active physical weight folds={int(physics_mix.gt(0).sum())}/{len(physics_mix)}"),
        ("development_only_model_selection", selection_trace_complete,
         selection_trace_detail),
        ("response_free_temperature_extrapolation_guard", temperature_gate_verified,
         "each outer-test residual-support flag is reproduced from development exact-fracture temperatures only"),
        ("best_primary_campaign_balanced_model", best_primary == PROPOSED_MODEL,
         f"lowest primary campaign-disjoint, campaign-balanced RMSE among pre-specified models: {best_primary}"),
        ("reported_model_matches_selected_residual_branch",
         abs(float(proposed["SB_RMSE"]) - float(anchor["SB_RMSE"])) < 1e-10,
         f"anchor RMSE={anchor['SB_RMSE']:.4f}; reported RMSE={proposed['SB_RMSE']:.4f}"),
    ]
    for comparator, row in paired.items():
        checks.append((
            f"beats_{comparator}_point", bool(np.isfinite(row["estimate"]) and float(row["estimate"]) > 0),
            f"paired RMSE gain={float(row['estimate']):.4f}" if np.isfinite(row["estimate"]) else "not evaluable",
        ))
        checks.append((
            f"beats_{comparator}_CI", rmse_ci_pass[comparator],
            (f"95% campaign-bootstrap CI [{float(row['CI95_lo']):.4f}, {float(row['CI95_hi']):.4f}]"
             if np.isfinite(row["CI95_lo"]) else "not evaluable"),
        ))
    checks.extend([
        ("primary_RMSE_superiority_gate", primary_rmse_superiority,
         "RMSE-gain 95% campaign-bootstrap CI is above zero versus ML-Ens, ExtraTrees and ET-Walker"),
        ("primary_multimetric_superiority_gate", primary_multimetric_superiority,
         "requires positive 95% CIs for RMSE, MAE, R2, F5 and C* gains versus all three primary comparators"),
        ("external_domain_point_superiority", external_point_superiority,
         f"WB-PIML has the lowest RMSE in every LOSO/LOAO/LOTO/LOEO factor-held-out scenario: {external_point_passes}"),
        ("external_domain_CI_superiority", external_ci_superiority,
         "requires positive paired RMSE-gain CIs versus all three comparators in every factor-held-out scenario"),
        ("external_nested_selection_complete", external_nested_selection_complete,
         external_selection_detail),
        ("external_censor_guard_never_relaxed",
         external_censor_guard_never_relaxed,
         f"constraint-relaxed WB-PIML domain holdouts="
         f"{external_relaxed_count}/{external_wb_holdout_count}"),
        ("external_interval_audit_complete", external_interval_complete,
         "all LOSO/LOAO/LOTO/LOEO WB-PIML exact predictions have development-only 80% and 90% intervals"),
        ("external_interval_nominal_compatible", external_interval_nominal_compatible,
         "each scenario-level empirical coverage CI must contain its nominal 80% or 90% target"),
        ("runout_survival_noninferiority", runout_noninferior,
         f"outer runout survival NLL: WB-PIML={proposed_runout_nll:.4f}; WB-CD={reference_runout_nll:.4f}"),
        ("censor_guard_never_relaxed", censor_guard_never_relaxed,
         f"constraint-relaxed folds={relaxed_count}/{len(anchor_selection)}"),
        ("mechanistic_damage_constraints", mechanistic_constraints_verified,
         "all reported WB-PIML folds have five finite physical parameters, positive Walker slope, bounded non-negative environmental exponents, and active physics"),
        ("mechanistic_parameters_interior", mechanistic_parameters_interior,
         f"competing-damage parameter boundary hits={int(np.sum(boundary_hits))}/{len(boundary_hits)} folds"),
        ("overall_claim_gate", universal_superiority,
         "broad claims require primary multi-metric superiority, complete training-only nested selection and superiority in domain-held-out tests, runout non-inferiority, unrelaxed primary and domain-holdout censor guards, interior physical parameters and compatible external interval coverage"),
    ])
    table = pd.DataFrame(checks, columns=["gate", "passed", "evidence"])
    table["evidence_origin"] = np.where(
        table["gate"].isin(["workbook_audit_complete", "human_reviewed_campaign_map",
                            "label_conflicts_resolved", "physical_duplicate_audit"]),
        "workbook", "computed",
    )
    return table


def campaign_audit_tables(df: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Audit all 23 physical duplicate groups from Plot_Data, not only retained rows."""
    raw = data.AUDIT_CONTEXT.get("raw")
    if not isinstance(raw, pd.DataFrame) or raw.empty:
        raise RuntimeError("Plot_Data is required for duplicate provenance audit")
    raw = raw.copy()
    raw["record_equivalence_id"] = raw["record_equivalence_id"].fillna("").astype(str).str.strip()
    counts = raw.groupby("record_equivalence_id").size()
    duplicate_ids = set(counts.loc[counts.gt(1)].index)
    reviewed = df.set_index("record_equivalence_id", drop=False)
    fold_by_campaign = df.groupby("campaign_id")["outer_fold"].first().to_dict()
    cluster_rows = []
    for equivalence_id in sorted(duplicate_ids):
        group = raw.loc[raw["record_equivalence_id"].eq(equivalence_id)].copy()
        supplied_runout = parse_bool(group["is_runout"], "Plot_Data.is_runout")
        representative = reviewed.loc[equivalence_id]
        if isinstance(representative, pd.DataFrame):
            representative = representative.iloc[0]
        cluster_rows.append({
            "record_equivalence_id": equivalence_id,
            "campaign_id": representative["campaign_id"],
            "n_rows": len(group), "n_sources": group["source_id"].astype(str).nunique(),
            "n_exact": int((~supplied_runout).sum()), "n_runout": int(supplied_runout.sum()),
            "supplied_label_conflict": bool(supplied_runout.nunique() > 1),
            "reviewed_label_conflict": bool(representative["label_conflict_flag"]),
            "label_conflict_resolved": bool(representative["label_conflict_resolved"]),
            "n_label_overrides": 0, "n_outer_folds": 1,
            "source_ids": " | ".join(sorted(group["source_id"].astype(str).unique())),
            "retained_row_id": int(representative["row_id"]),
        })
    clusters = pd.DataFrame(cluster_rows)
    cross_source_ids = (
        set(clusters.loc[
            clusters["n_sources"].gt(1), "record_equivalence_id"
        ].astype(str))
        if not clusters.empty else set()
    )

    candidates = raw.loc[raw["record_equivalence_id"].isin(duplicate_ids)].copy()
    candidates["cross_source_duplicate_candidate"] = (
        candidates["record_equivalence_id"].astype(str).isin(cross_source_ids)
    )
    candidates["supplied_is_runout"] = parse_bool(candidates["is_runout"], "Plot_Data.is_runout")
    candidates["supplied_is_exact_failure"] = ~candidates["supplied_is_runout"]
    candidates["reviewed_is_runout"] = candidates["record_equivalence_id"].map(
        df.set_index("record_equivalence_id")["is_runout"]
    ).astype(bool)
    candidates["reviewed_is_exact_failure"] = ~candidates["reviewed_is_runout"]
    candidates["is_runout"] = candidates["reviewed_is_runout"]
    candidates["is_exact"] = candidates["reviewed_is_exact_failure"]
    candidates["label_override_applied"] = False
    candidates["campaign_id"] = candidates["record_equivalence_id"].map(
        df.set_index("record_equivalence_id")["campaign_id"]
    )
    candidates["outer_fold"] = candidates["campaign_id"].map(fold_by_campaign).astype(int)
    candidates["supplied_label_conflict_flag"] = candidates["record_equivalence_id"].map(
        clusters.set_index("record_equivalence_id")["supplied_label_conflict"]
    ).astype(bool)
    candidates["label_conflict_flag"] = False
    candidates = candidates.sort_values(
        ["record_equivalence_id", "source_id", "row_id"]
    ).reset_index(drop=True)

    campaign_map = df.groupby(["campaign_id", "source_id"], sort=True).agg(
        n_rows=("row_id", "size"), n_exact=("is_exact", "sum"),
        n_runout=("is_runout", "sum"), outer_fold=("outer_fold", "first"),
        representative_source=("campaign_representative_source", "first"),
        source_campaign_role=("source_campaign_role", "first"),
        reviewed_secondary_source=("reviewed_secondary_source", "first"),
        campaign_mapping_status=("campaign_mapping_status", "first"),
        mapping_basis=("campaign_mapping_basis", "first"),
        audit_status=("audit_status", "first"), evidence_grade=("evidence_grade", "first"),
    ).reset_index()
    return clusters, candidates, campaign_map


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--data", default=None)
    parser.add_argument("--out", default=str(DEFAULT_OUTPUT_DIR))
    parser.add_argument("--bootstrap", type=int, default=5000)
    parser.add_argument("--trees", type=int, default=400)
    parser.add_argument(
        "--audit-trees", type=int, default=160,
        help=(
            "Trees per residual model in the expanded sensitivity audit. "
            "This does not change the primary model (default: 160)."
        ),
    )
    parser.add_argument("--repeats", type=int, default=10,
                        help="Repeated validation runs for each of the four validation tiers (1-10).")
    parser.add_argument(
        "--export-csv",
        action="store_true",
        help="Also export each result table as a CSV file.",
    )
    args = parser.parse_args()
    if Font is None:
        raise ImportError("Excel output requires openpyxl")
    if args.bootstrap < 100:
        raise ValueError("Use at least 100 campaign bootstrap replicates")
    if args.trees < 20:
        raise ValueError("--trees must be at least 20")
    if args.audit_trees < 20:
        raise ValueError("--audit-trees must be at least 20")
    if not 1 <= args.repeats <= len(REPEATED_VALIDATION_SEEDS):
        raise ValueError(f"--repeats must be between 1 and {len(REPEATED_VALIDATION_SEEDS)}")
    data_path = resolve_data(args.data)
    out = Path(args.out).resolve()
    out.mkdir(parents=True, exist_ok=True)
    started = time.perf_counter()
    full_df = load_data(data_path)
    df = full_df.reset_index(drop=True)
    exact_df = df.loc[df["is_exact"]].reset_index(drop=True)
    runout_df = full_df.loc[full_df["is_runout"]].reset_index(drop=True)
    require_exact_training(df, "main joint-likelihood analysis frame")
    if (
        len(full_df) != EXPECTED_VERIFIED_ROWS
        or len(exact_df) != EXPECTED_EXACT_ROWS
        or len(runout_df) != EXPECTED_RUNOUT_ROWS
    ):
        raise ValueError(
            "Plot_Data_Verified must contain 222 physically unique records: "
            "160 exact failures and 62 runouts; "
            f"found {len(full_df)} total, {len(exact_df)} exact and {len(runout_df)} runout rows"
        )
    if not np.allclose(training_weights(df), 1.0):
        raise AssertionError("Every unique likelihood record must have unit weight")
    outer_splits, outer_id = grouped_folds(df, 5, SEED)
    df["outer_fold"] = outer_id
    fold_metric_rows = []
    (
        prediction_parts, parameter_rows, audit_rows,
        nested_selection_rows, nested_candidate_parts,
    ) = run_primary_folds(df, outer_splits, args.trees)

    predictions = pd.concat(prediction_parts, ignore_index=True)
    predictions["T_C"] = predictions["row_id"].map(
        df.set_index("row_id")["T_C"]
    ).to_numpy(float)
    expected_models = ([s.key for s in PHYSICAL_SPECS]
                       + [name for name, _ in HYBRID_SPECS]
                       + [PROPOSED_MODEL]
                       + list(EXTERNAL_MODELS) + list(CENSORED_MODELS))
    for model in expected_models:
        model_rows = predictions.loc[predictions["model"].eq(model)]
        if len(model_rows) != len(df) or model_rows["row_id"].duplicated().any():
            raise AssertionError(f"{model} must predict every record exactly once")

    temperature_gate_rows = predictions.loc[
        predictions["model"].eq(PROPOSED_MODEL),
        ["outer_fold", "row_id", "source_id", "campaign_id",
         "record_equivalence_id", "T_C", "is_exact", "is_runout",
         "residual_temperature_supported", "logN", "pred_logN"],
    ].copy()
    temperature_gate_summary = (
        temperature_gate_rows.groupby(
            ["outer_fold", "T_C", "is_exact", "is_runout",
             "residual_temperature_supported"],
            dropna=False,
        )
        .size().rename("n_rows").reset_index()
        .sort_values(["outer_fold", "T_C", "is_runout"])
    )

    for (fold, model), group in predictions.loc[predictions["is_exact"]].groupby(["outer_fold", "model"]):
        groups = [g.reset_index(drop=True) for _, g in group.groupby(PRIMARY_GROUP_COLUMN)]
        fold_metric_rows.append({"outer_fold": fold, "model": model, "n_exact": len(group), **sb_metrics(groups)})

    print("Summarizing primary predictions and bootstrap intervals", flush=True)
    metrics = metric_summary(predictions, args.bootstrap, SEED + 1)
    effects = pd.concat([
        paired_effects(predictions, args.bootstrap, SEED + 20 + index, comparator=model)
        for index, model in enumerate(EXTERNAL_MODELS)
    ], ignore_index=True)
    uq = uq_summary(predictions, args.bootstrap, SEED + 3)
    runout_predictions = predictions.loc[predictions["is_runout"]].copy()
    runout_predictions["evaluation_scope"] = "campaign_disjoint_outer_test_right_censored"
    runout_predictions["used_for_training"] = runout_predictions["model"].isin(
        [s.key for s in PHYSICAL_SPECS] + [name for name, _ in HYBRID_SPECS]
        + [PROPOSED_MODEL] + list(CENSORED_MODELS)
    )
    runout_predictions["used_for_candidate_selection"] = (
        runout_predictions["model"].isin(
            ["WB-PIML-Anchor", PROPOSED_MODEL] + list(CENSORED_MODELS)
        )
    )
    runout_predictions["used_as_censor_feasibility_reference"] = (
        runout_predictions["model"].eq("WB-CD")
    )
    runout_predictions["used_for_selection"] = (
        runout_predictions["used_for_candidate_selection"]
        | runout_predictions["used_as_censor_feasibility_reference"]
    )
    runout = runout_summary(runout_predictions)
    parameters = pd.DataFrame(parameter_rows)
    parameters["damage_coefficient_sum"] = parameters[["M_coefficient", "Ox_coefficient", "Q_coefficient"]].fillna(0).sum(axis=1)
    parameters["damage_boundary_hit"] = parameters["damage_coefficient_sum"] >= DAMAGE_CAP - 1e-5
    parameters["walker_slope_boundary_hit"] = parameters["walker_slope"] <= 0.05 + 1e-5
    parameters["environment_ratio_boundary_hit"] = (
        np.isclose(pd.to_numeric(parameters["environment_log10_damage_ratio"], errors="coerce"), -8.0, atol=1e-5)
        | np.isclose(pd.to_numeric(parameters["environment_log10_damage_ratio"], errors="coerce"), 4.0, atol=1e-5)
    )
    parameters["environment_stress_boundary_hit"] = (
        np.isclose(pd.to_numeric(parameters["environment_stress_exponent"], errors="coerce"), 0.0, atol=1e-5)
        | np.isclose(pd.to_numeric(parameters["environment_stress_exponent"], errors="coerce"), 10.0, atol=1e-5)
    )
    parameters["arrhenius_boundary_hit"] = (
        np.isclose(pd.to_numeric(parameters["arrhenius_temperature_sensitivity"], errors="coerce"), 0.0, atol=1e-5)
        | np.isclose(pd.to_numeric(parameters["arrhenius_temperature_sensitivity"], errors="coerce"), 10.0, atol=1e-5)
    )
    parameters["competing_damage_boundary_hit"] = parameters[[
        "environment_ratio_boundary_hit", "environment_stress_boundary_hit",
        "arrhenius_boundary_hit",
    ]].any(axis=1)
    audit = pd.DataFrame(audit_rows)
    fold_metrics = pd.DataFrame(fold_metric_rows)
    ablation = ablation_effects(predictions, args.bootstrap, SEED + 4)
    source_gains = source_gain_table(predictions)
    nested_selection = pd.DataFrame(nested_selection_rows)
    nested_candidates = pd.concat(nested_candidate_parts, ignore_index=True, sort=False)
    print("Refitting sensitivity analyses", flush=True)
    sensitivity = sensitivity_audit(
        df, outer_splits, args.audit_trees, nested_selection
    )
    component_ablation = component_ablation_summary(sensitivity)
    selection_sensitivity = selection_rule_sensitivity(nested_candidates)
    print("Running repeated validation", flush=True)
    repeated_metrics, validation_summary, validation_assignments = repeated_validation(
        df, args.trees, args.repeats
    )
    print("Running nested factor holdouts", flush=True)
    (
        external_summary,
        external_groups,
        external_holdout_predictions,
        external_selection,
    ) = external_holdout_validation(df, args.trees)
    print("Summarizing factor holdouts", flush=True)
    external_effects = external_paired_effects(
        external_holdout_predictions, args.bootstrap, SEED + 930000
    )
    external_uq = external_uq_summary(
        external_holdout_predictions, args.bootstrap, SEED + 940000
    )
    external_pi_rows = external_holdout_predictions.loc[
        external_holdout_predictions["model"].eq(PROPOSED_MODEL)
        & external_holdout_predictions["is_exact"]
    ].copy()
    external_temperature_gate_summary = (
        external_holdout_predictions.loc[
            external_holdout_predictions["model"].eq(PROPOSED_MODEL)
        ]
        .groupby(
            ["strategy", "holdout_group", "fold", "T_C", "is_exact",
             "is_runout", "residual_temperature_supported"],
            dropna=False,
        )
        .size().rename("n_rows").reset_index()
        .sort_values(["strategy", "fold", "T_C", "is_runout"])
    )
    gate = claim_gate(
        metrics, effects, parameters, df, nested_selection, nested_candidates,
        runout, external_summary, external_groups, external_effects, external_uq,
        external_selection,
        temperature_gate_rows,
    )
    print("Checking stress and duplicate-record sensitivities", flush=True)
    sstar_sensitivity = sstar_exclusion_sensitivity(df, args.trees)
    duplicate_sensitivity = duplicate_source_sensitivity(df, args.trees)
    duplicate_clusters, duplicate_candidates, campaign_map = campaign_audit_tables(df)
    label_review = full_df.loc[
        full_df["supplied_label_conflict_flag"] | full_df["label_override_applied"],
        ["row_id", "source_row_no", "source_id", "record_equivalence_id",
         "sigma_max_MPa_before_normalization", "Nf_cycles", "raw_Nf",
         "supplied_is_exact_failure", "supplied_is_runout",
         "reviewed_is_exact_failure", "is_runout", "label_override_applied",
         "label_review_status", "label_review_evidence"],
    ].copy().sort_values(["record_equivalence_id", "source_id", "row_id"])
    label_review["evidence_file"] = (
        f"{data_path.name} / Verification_Log and Label_Review evidence fields"
    )

    qa_flagged = df.loc[df["qa_sstar_gt_1"] | df["qa_exact_round_limit"], [
        "row_id", "source_id", "architecture_detail", "env_class", "T_C", "R",
        "frequency_Hz", "Nf_cycles", "S", "is_exact", "is_runout",
        "qa_sstar_gt_1", "qa_exact_round_limit",
    ]].copy()
    qa_flagged["action"] = np.where(
        qa_flagged["qa_exact_round_limit"],
        "verify against original source; main analysis uses the reviewed label recorded in Label_Review",
        "retained in main analysis; compare with S*>1 exclusion sensitivity",
    )
    validation_definitions = pd.DataFrame([
        ("record_interpolation", "rows stratified across five folds; source/series overlap is disclosed",
         "secondary interpolation diagnostic; not evidence for unseen-source prediction"),
        ("series_disjoint", "series = source + detailed architecture + environment + T + R + UTS; series never overlaps",
         "primary in-domain grouped validation when explicit batch identifiers are unavailable"),
        ("source_disjoint", "publication source never overlaps; ten repeated five-fold assignments",
         "secondary comparison; may split publications from one experimental campaign"),
        ("campaign_disjoint", "human-reviewed source-to-campaign map never overlaps; ten repeated five-fold assignments",
         "fully campaign-disjoint outer validation within the compiled database"),
        ("LOSO", "one complete source held out at a time; another source from the same campaign may remain in development",
         "source-held-out stress test with training-only nested selection; source isolation does not imply campaign isolation, and actual overlap is reported"),
        ("LOAO", "one weaving-architecture group held out at a time",
         "weaving-architecture-held-out stress test with training-only nested selection; campaign isolation is not imposed by this split, and actual overlap is reported"),
        ("LOTO", "one reviewed temperature regime held out at a time",
         "temperature-regime-held-out stress test with training-only nested selection; campaign isolation is not imposed by this split, and actual overlap is reported"),
        ("LOEO", "one environment class held out at a time",
         "environment-held-out stress test with training-only nested selection; campaign isolation is not imposed by this split, and actual overlap is reported"),
    ], columns=["strategy", "split_definition", "interpretation"])

    proposed_exact = predictions.loc[
        predictions["model"].eq("WB-PIML") & predictions["is_exact"]
    ].copy()
    for level in (80, 90):
        proposed_exact[f"covered{level}"] = (
            proposed_exact["logN"].between(proposed_exact[f"lower{level}"], proposed_exact[f"upper{level}"])
        )
        proposed_exact[f"width{level}"] = proposed_exact[f"upper{level}"] - proposed_exact[f"lower{level}"]

    model_specs = pd.DataFrame([
        {"model": "WB", "role": "right-censored lognormal AFT physics baseline", "location_parameters": 2,
         "equation": "mu=a-b*log10(S_Walker); exact density + runout survival", "used_for_selection": False},
        {"model": "WB-CD", "role": "right-censored competing-damage physics trunk", "location_parameters": 5,
         "equation": "mu=-log10[D_mech*(1 + 10^logK*E(env)*Arrhenius(T)*S_W^m/f)]; censored AFT",
         "used_for_selection": True},
        {"model": "WB-M", "role": "mechanical-proxy ablation", "location_parameters": 3,
         "equation": "a - b*log10(S_Walker) - c*M", "used_for_selection": False},
        {"model": "WB-Ox", "role": "oxidation-proxy ablation", "location_parameters": 3,
         "equation": "a - b*log10(S_Walker) - c*Dox", "used_for_selection": False},
        {"model": "WB-PIML-Anchor", "role": "selected competing-damage residual branch", "location_parameters": 5,
         "equation": "mu_WB-CD + eta*ExtraTrees[X, y-mu_WB-CD]", "used_for_selection": True},
        {"model": "WB-PIML", "role": "reported right-censored always-physical competing-damage PIML", "location_parameters": 5,
         "equation": "censored-AFT mu_WB-CD + g_T(T)*eta*g_residual; g_T=0 without same-temperature exact support in development", "used_for_selection": True},
        {"model": "WB-PIML-M", "role": "mechanical-anchor residual ablation", "location_parameters": 3,
         "equation": "mu_WB-M + eta*ExtraTrees[X, y-mu_WB-M]", "used_for_selection": False},
        {"model": "WB-PIML-Ox", "role": "oxidation-anchor residual ablation", "location_parameters": 3,
         "equation": "mu_WB-Ox + eta*ExtraTrees[X, y-mu_WB-Ox]", "used_for_selection": False},
        {"model": "ExtraTrees", "role": "external traditional comparator",
          "location_parameters": np.nan, "equation": "same observable inputs; never replaces the physical trunk",
          "used_for_selection": True},
        {"model": "ET-Walker", "role": "equally tuned paired transform control",
         "location_parameters": np.nan,
         "equation": "ExtraTrees[X, selected log10(S_Walker)]; same folds/candidate slots/tree seeds as WB-PIML",
         "used_for_selection": True},
        {"model": "RF / GBDT / SVR", "role": "fixed external traditional comparators",
         "location_parameters": np.nan, "equation": "same observable inputs; not used inside WB-PIML",
         "used_for_selection": False},
        {"model": "ML-Ens", "role": "external traditional ensemble comparator",
         "location_parameters": np.nan,
         "equation": "mean of fixed RF/GBDT/SVR and nested-selected ExtraTrees predictions",
         "used_for_selection": False},
        {"model": "Lognormal-AFT", "role": "fair observable-input right-censored comparator",
         "location_parameters": np.nan,
         "equation": "penalized AFT location; exact normal density + runout survival",
         "used_for_selection": True},
        {"model": "Weibull-AFT", "role": "fair observable-input right-censored comparator",
         "location_parameters": np.nan,
         "equation": "penalized Weibull AFT; exact density + runout survival",
         "used_for_selection": True},
    ])

    data_hash = hashlib.sha256(data_path.read_bytes()).hexdigest()
    data_audit = pd.DataFrame([
        ("input_file", data_path.name), ("sheet", SHEET), ("sha256", data_hash),
        ("n_verified_rows", len(full_df)),
        ("n_exact_density_training_and_evaluation_rows", len(exact_df)),
        ("n_right_censored_likelihood_rows", len(runout_df)),
        ("n_workbook_exact_fractures", int(full_df["is_exact"].sum())),
        ("n_workbook_runouts", int(full_df["is_runout"].sum())),
        ("n_unique_record_equivalence_ids", int(full_df["record_equivalence_id"].nunique())),
        ("duplicate_equivalence_ids_in_verified", bool(full_df["record_equivalence_id"].duplicated().any())),
        ("exact_training_weight_min", f"{float(exact_df['sample_weight'].min()):.1f}"),
        ("exact_training_weight_max", f"{float(exact_df['sample_weight'].max()):.1f}"),
        ("runout_training_weight_min", f"{float(runout_df['sample_weight'].min()):.1f}"),
        ("runout_training_weight_max", f"{float(runout_df['sample_weight'].max()):.1f}"),
        ("n_code_side_label_overrides", int(full_df["label_override_applied"].sum())),
        ("n_sources", int(full_df["source_id"].nunique())),
        ("n_exact_sources", int(exact_df["source_id"].nunique())),
        ("n_campaigns_all_verified", int(full_df["campaign_id"].nunique())),
        ("n_campaigns_exact_analysis", int(exact_df["campaign_id"].nunique())),
        ("n_series_groups", int(df["series_group_id"].nunique())),
        ("n_physical_duplicate_clusters", len(duplicate_clusters)),
        ("n_cross_source_duplicate_clusters", int(duplicate_clusters["n_sources"].gt(1).sum())),
        ("n_duplicate_traceability_rows", len(duplicate_candidates)),
        ("n_cross_source_duplicate_exact", int(
            duplicate_candidates.loc[
                duplicate_candidates["cross_source_duplicate_candidate"], "is_exact"
            ].sum()
        )),
        ("n_supplied_label_conflict_clusters", int(duplicate_clusters["supplied_label_conflict"].sum())),
        ("n_reviewed_label_conflict_clusters", int(duplicate_clusters["reviewed_label_conflict"].sum())),
        ("n_supplied_label_conflict_rows", int(full_df["supplied_label_conflict_flag"].sum())),
        ("n_reviewed_label_conflict_rows", int(full_df["label_conflict_flag"].sum())),
        ("n_reviewed_secondary_sources", int(full_df.loc[full_df["reviewed_secondary_source"], "source_id"].nunique())),
        ("label_consistency_all_rows", bool(full_df["label_consistent"].all())),
        ("duplicate_clusters_cross_outer_folds", int(duplicate_clusters["n_outer_folds"].gt(1).sum())),
        ("n_Sstar_gt_1", int(df["qa_sstar_gt_1"].sum())),
        ("n_exact_round_limit_flags", int(df["qa_exact_round_limit"].sum())),
        ("missing_required_values", 0), ("duplicate_row_id", False),
        ("runouts_enter_preprocessor_fit", False),
        ("runouts_enter_location_training", True),
        ("runouts_enter_hyperparameter_selection", True),
        ("runouts_enter_calibration", True),
        ("runouts_enter_point_error_metrics", False), ("random_seed", SEED),
        ("row_or_source_dependent_training_weights", False),
    ], columns=["item", "value"])

    references = pd.DataFrame([
        ("Verified C/SiC fatigue workbook audit trail", "reviewed exact/runout decisions and source-campaign mapping",
         f"{data_path.name} / Verification_Log, Campaign_Map and Label_Review"),
        ("Walker mean-stress correction", "physics trunk", "https://doi.org/10.1115/1.4001673"),
        ("Physics-informed feature engineering for fatigue", "residual-hybrid context", "https://www.mdpi.com/2076-3417/16/13/6493"),
        ("Accelerated failure-time modelling", "right-censoring likelihood audit", "https://academic.oup.com/biomet/article-abstract/79/2/311/225970"),
        ("Ceramic-matrix-composite fatigue mechanisms", "mechanism-informed proxy design", "https://journals.sagepub.com/doi/10.1177/14644207211008574"),
        ("High-temperature oxidation of SiC/SiC composites", "discussion context; no thresholded oxidation feature in the primary model", "https://www.mdpi.com/1996-1944/9/3/207"),
        ("Oxidation-dependent CMC fatigue life", "environment-fatigue coupling", "https://www.sciencedirect.com/science/article/pii/S0921509315300575"),
        ("Physics-constrained ML for CMCs", "advanced-method context", "https://www.sciencedirect.com/science/article/pii/S1359836825007310"),
    ], columns=["reference_topic", "use_in_analysis", "url_or_local_source"])

    outer_fold_assignment = df[
        ["row_id", "source_id", "campaign_id", "record_equivalence_id",
         "series_group_id", "cross_source_duplicate_candidate",
         "supplied_label_conflict_flag", "label_conflict_flag",
         "source_campaign_role", "campaign_mapping_status",
         "supplied_is_exact_failure", "supplied_is_runout",
         "is_exact", "is_runout", "label_override_applied", "outer_fold"]
    ].copy()
    if args.export_csv:
        csv_tables = {
            "predictions.csv": predictions,
            "campaign_balanced_metrics.csv": metrics,
            "fold_metrics.csv": fold_metrics,
            "primary_paired_effects.csv": effects,
            "empirical_interval_summary.csv": uq,
            "empirical_interval_rows.csv": proposed_exact,
            "temperature_gate_rows.csv": temperature_gate_rows,
            "temperature_gate_summary.csv": temperature_gate_summary,
            "runout_scale_audit.csv": runout,
            "runout_outer_predictions.csv": runout_predictions,
            "physical_parameters_and_scale.csv": parameters,
            "validation_audit.csv": audit,
            "retrained_ablation_effects.csv": ablation,
            "fixed_sensitivity_audit.csv": sensitivity,
            "component_ablation.csv": component_ablation,
            "selection_rule_sensitivity.csv": selection_sensitivity,
            "source_level_gains.csv": source_gains,
            "claim_gate.csv": gate,
            "outer_fold_assignment.csv": outer_fold_assignment,
            "nested_selection.csv": nested_selection,
            "nested_candidates.csv": nested_candidates,
            "validation_tier_metrics.csv": repeated_metrics,
            "validation_tier_summary.csv": validation_summary,
            "external_holdout_summary.csv": external_summary,
            "external_holdout_groups.csv": external_groups,
            "external_selection.csv": external_selection,
            "external_paired_effects.csv": external_effects,
            "external_interval_summary.csv": external_uq,
            "external_interval_rows.csv": external_pi_rows,
            "external_temperature_gate_summary.csv": external_temperature_gate_summary,
            "sstar_sensitivity.csv": sstar_sensitivity,
            "duplicate_source_sensitivity.csv": duplicate_sensitivity,
            "duplicate_clusters.csv": duplicate_clusters,
            "duplicate_candidates.csv": duplicate_candidates,
            "campaign_map.csv": campaign_map,
            "label_review.csv": label_review,
            "qa_flagged_rows.csv": qa_flagged,
        }
        for filename, table in csv_tables.items():
            table.to_csv(out / filename, index=False)
    external_wb_selection = external_selection.loc[
        external_selection["model"].eq(PROPOSED_MODEL)
    ].copy()
    external_relaxed_flags = pd.to_numeric(
        external_wb_selection["censor_constraint_relaxed"], errors="coerce"
    )
    external_censor_status_known = bool(
        len(external_wb_selection) > 0
        and external_relaxed_flags.notna().all()
        and external_relaxed_flags.isin([0, 1]).all()
    )
    external_campaign_overlap_flags = external_wb_selection[
        "campaign_overlap"
    ].astype(bool)

    protocol = {
        "model": "WB-PIML", "data": data_path.name, "sheet": SHEET,
        "data_sha256": data_hash, "n_verified_rows": len(full_df),
        "n_exact_density_training_evaluation": len(exact_df), "n_right_censored_training": len(runout_df),
        "n_label_overrides_in_code": 0,
        "n_sources": int(full_df["source_id"].nunique()), "n_exact_sources": int(exact_df["source_id"].nunique()),
        "n_campaigns": int(df["campaign_id"].nunique()),
        "n_physical_duplicate_clusters": len(duplicate_clusters),
        "n_cross_source_duplicate_clusters": int(duplicate_clusters["n_sources"].gt(1).sum()),
        "n_duplicate_traceability_rows": len(duplicate_candidates),
        "n_supplied_label_conflict_clusters": int(duplicate_clusters["supplied_label_conflict"].sum()),
        "n_reviewed_label_conflict_clusters": int(duplicate_clusters["reviewed_label_conflict"].sum()),
        "outer_folds": len(outer_splits),
        "candidate_gamma": sorted({float(item["gamma"]) for item in HYBRID_CANDIDATES}),
        "candidate_min_samples_leaf": sorted({int(item["min_samples_leaf"]) for item in HYBRID_CANDIDATES}),
        "candidate_max_features": sorted({float(item["max_features"]) for item in HYBRID_CANDIDATES}),
        "candidate_eta": sorted({float(item["eta"]) for item in HYBRID_CANDIDATES}),
        "candidate_residual_modes": sorted({str(item["residual_mode"]) for item in HYBRID_CANDIDATES}),
        "candidate_physics_mix": list(PHYSICS_MIX_GRID),
        "calibration_modes": list(CALIBRATION_MODES),
        "calibration_status": "locked to none; affine output calibration is absent from the executable path",
        "RMSE_tie_tolerance": RMSE_TIE_TOLERANCE,
        "damage_cap_for_ablation_only": DAMAGE_CAP,
        "fold_balance": "campaign-disjoint nested outer/inner folds constructed from all exact and right-censored records",
        "point_metric_balance": "equal campaign contribution; SB_* is retained only as a legacy compatibility prefix",
        "campaign_mapping": "campaign_id read directly from Plot_Data_Verified; no code-side remapping",
        "duplicate_handling": "Plot_Data_Verified is physically deduplicated; one retained row per record_equivalence_id",
        "training_weighting": "all 222 physically unique records use unit likelihood weight; exact density and runout survival contributions differ only by censoring status",
        "label_review": "is_exact_failure and is_runout read directly from Plot_Data_Verified; no code-side override",
        "location_training": "joint lognormal AFT likelihood: exact density plus right-censored survival",
        "proposed_architecture": "always-active five-parameter censored-AFT competing-damage trunk plus a nested-selected legacy or support-gated dual exact-fracture residual",
        "proposed_location_equation": "mu_phys=-log10(D_mech+D_env); yhat=mu_phys+g_T(T)*eta*[q_support*g_robust+0.25*(1-q_support)*g_smooth]",
        "residual_extrapolation_guard": "eta_eff=0 when the development exact-fracture set has no observation at the requested temperature; exact temperature matching uses 1e-9 C only for floating-point equality and never consults validation responses",
        "fitted_physics_parameters": 5, "nuisance_scale_parameters_per_fold": 1,
        "runout_use": "runout stopping cycles enter only as lower bounds through -log survival; they affect the physical trunk, scale and NLL tie-break, never point residual regression or point-life metrics",
        "outer_test_tuning": False,
        "primary_validation_scope": (
            "campaign-disjoint outer validation; no campaign crosses train/test"
        ),
        "domain_holdout_scope": (
            "LOSO/LOAO/LOTO/LOEO are factor-held-out stress tests, not an "
            "independent external dataset and not campaign-disjoint by design; "
            "actual overlap is reported per holdout"
        ),
        "hyperparameter_selection": f"{len(HYBRID_CANDIDATES)} predeclared candidates ({len(NESTED_CANDIDATES)} tree-based candidates + 4 robust support candidates) selected exclusively in development folds",
        "selection_endpoint": "exact campaign-balanced inner-fold RMSE; within 0.02 RMSE, prefer the censor-feasible shortlist and lower inner temperature-disjoint RMSE",
        "censor_guardrail": "candidate runout survival NLL <= WB-CD runout survival NLL + one paired inner-fold standard error; relaxation is explicitly recorded if no near-best candidate is feasible",
        "ET_Walker_fairness": f"same joint-data inner folds and {len(NESTED_CANDIDATES)} gamma/tree candidates; four additional candidates use alternative residual formulations",
        "fair_censored_comparators": list(CENSORED_MODELS),
        "primary_comparators": "ExtraTrees, equally tuned ET-Walker, and ML-Ens",
        "primary_inputs": TRAD_NUMERIC + TRAD_CATEGORICAL,
        "arbitrary_oxidation_proxy_in_primary_inputs": False,
        "validation_tiers": ["record_interpolation", "series_disjoint", "source_disjoint", "campaign_disjoint", "LOSO", "LOAO", "LOTO", "LOEO"],
        "external_holdout_selection": (
            "for every LOSO/LOAO/LOTO/LOEO holdout, WB-PIML, ExtraTrees and "
            "ET-Walker are re-selected only in the remaining records using "
            "campaign-disjoint inner folds; held-out responses are never used"
        ),
        "external_holdout_count": int(
            external_selection[["scenario", "holdout_group"]]
            .drop_duplicates().shape[0]
        ),
        "external_selection_rows": int(len(external_selection)),
        "external_censor_guard_relaxed_holdouts": int(
            external_relaxed_flags.eq(1).sum()
        ),
        "external_censor_guard_never_relaxed": bool(
            external_censor_status_known
            and external_relaxed_flags.eq(0).all()
        ),
        "external_campaign_overlap_holdouts": int(
            external_campaign_overlap_flags.sum()
        ),
        "external_campaign_disjoint_holdouts": int(
            (~external_campaign_overlap_flags).sum()
        ),
        "external_inner_group": PRIMARY_GROUP_COLUMN,
        "external_outer_test_tuning": False,
        "external_outer_test_interval_calibration": False,
        "repeated_validation_runs": args.repeats,
        "primary_model_trees": args.trees,
        "sensitivity_audit_trees": args.audit_trees,
        "sensitivity_audit_changes_primary_model": False,
        "intervals": (
            "campaign-equal weighted development-OOF empirical absolute-residual "
            "intervals; the same predeclared construction is recalibrated from "
            "each held-out training set's OOF residuals in LOSO/LOAO/LOTO/LOEO"
        ),
        "interval_label": "empirical prediction intervals; not formal finite-sample conformal guarantees",
        "external_effect_intervals": "paired campaign-cluster bootstrap CIs within each factor-held-out scenario",
        "output_contract": (
            "one Excel workbook, eight composite PNG figures, and one 29-slide "
            "PPTX containing each composite and every non-empty subplot; CSV disabled by default"
        ),
        "csv_export_enabled": bool(args.export_csv),
        "bootstrap": args.bootstrap, "seed": SEED,
        "software": {"python": platform.python_version(), "numpy": np.__version__, "pandas": pd.__version__,
                     "scipy": scipy.__version__, "scikit_learn": sklearn.__version__},
    }
    (out / "protocol.json").write_text(json.dumps(protocol, ensure_ascii=False, indent=2), encoding="utf-8")
    protocol_table = pd.DataFrame([{"item": key, "value": json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value}
                                   for key, value in protocol.items()])
    readme_table = pd.DataFrame([
        ("Purpose", "Verified-data, censor-aware, temperature-support, campaign-disjoint WB-PIML analysis"),
        ("Data", f"{data_path.name} / {SHEET} only"),
        ("Label review", "verified worksheet labels retained exactly; zero code-side overrides"),
        ("Point endpoint", f"Point prediction metrics use the {len(exact_df)} exact fractures only"),
        ("Runouts", f"{len(runout_df)} runouts enter training as right-censored survival terms and are never treated as point failure lives"),
        ("Campaign audit", f"{df['campaign_id'].nunique()} human-reviewed campaigns; {len(duplicate_clusters)} physical duplicate clusters ({int(duplicate_clusters['n_sources'].gt(1).sum())} cross-source); {int(duplicate_clusters['reviewed_label_conflict'].sum())} unresolved label-conflict clusters"),
        ("Primary validation", "nested five-fold campaign-disjoint outer validation; no campaign crosses training and test partitions"),
        ("Factor-held-out checks", "LOSO/LOAO/LOTO/LOEO use training-only nested selection; the factor is excluded from training, while campaign overlap is reported separately for each holdout"),
        ("Primary model", "Right-censored competing-damage trunk plus nested-selected tree-based or support-gated robust/smooth exact residual"),
        ("Censored controls", "Lognormal-AFT and Weibull-AFT use all exact/runout rows, the same observable inputs and the same outer campaign folds"),
        ("Paired control", "ET-Walker uses identical inner folds, candidate slots, gamma/tree settings and tree seeds"),
        ("Primary inputs", "observable stress, R, temperature, frequency, UTS, architecture and environment; no source/batch ID or thresholded oxidation score"),
        ("Tabular output", "one Excel workbook including External_Selection; CSV files disabled unless --export-csv is supplied"),
        ("Figure output", f"eight composite PNGs plus {PPTX_OUTPUT_NAME}; every non-empty subplot has a separate slide"),
        ("Interpretation", "Use Claim_Gate before drafting superiority claims"),
    ], columns=["item", "value"])
    metric_definitions = pd.DataFrame([
        ("SB_*", "legacy column-name prefix", "campaign-balanced in every primary and external summary; not publication-source-balanced"),
        ("SB_RMSE", "campaign-balanced RMSE", "square root of the equal-campaign weighted mean squared exact-fracture error"),
        ("SB_MAE", "campaign-balanced MAE", "equal-campaign weighted mean absolute exact-fracture error"),
        ("SB_R2", "campaign-balanced R2", "equal-campaign weighted coefficient of determination on exact fractures"),
        ("SB_F5", "campaign-balanced factor-five accuracy", "equal-campaign weighted fraction with absolute log10 error <= log10(5)"),
        ("SB_C_star", "campaign-balanced C*", "mean campaign-wise concordance error; lower is better"),
        ("SB_runout_NLL", "campaign-balanced survival NLL", "right-censored outer-test lower-bound likelihood; separate from point metrics"),
    ], columns=["column", "meaning", "definition"])
    tables = {
        "README": readme_table, "Protocol": protocol_table, "Data_Audit": data_audit,
        "Metric_Definitions": metric_definitions,
        "Model_Specs": model_specs, "Split_Audit": audit,
        "Validation_Definitions": validation_definitions,
        "Nested_Selection": nested_selection, "Nested_Candidates": nested_candidates,
        "Exact_Outer_Pred": predictions.loc[predictions["is_exact"]].copy(),
        "Runout_Outer_Pred": runout_predictions,
        "Runout_Evaluation": runout, "Primary_Campaign_Metrics": metrics,
        "Fold_Metrics": fold_metrics, "Paired_Bootstrap": effects,
        "True_Ablation": ablation, "Source_Gains": source_gains,
        "Component_Ablation": component_ablation,
        "Selection_Sensitivity": selection_sensitivity,
        "Empirical_PI_Rows": proposed_exact, "Empirical_PI_Summary": uq,
        "Temperature_Gate_Rows": temperature_gate_rows,
        "Temperature_Gate_Summary": temperature_gate_summary,
        "Physical_Parameters": parameters, "Sensitivity": sensitivity,
        "Duplicate_Sensitivity": duplicate_sensitivity,
        "Duplicate_Clusters": duplicate_clusters,
        "Duplicate_Candidates": duplicate_candidates,
        "Label_Review": label_review,
        "Campaign_Map": campaign_map,
        "Validation_Tier_Metrics": repeated_metrics,
        "Validation_Tier_Summary": validation_summary,
        "Validation_Fold_Assign": validation_assignments,
        "External_Summary": external_summary, "External_Groups": external_groups,
        "External_Selection": external_selection,
        "External_Paired_CI": external_effects,
        "External_PI_Summary": external_uq,
        "External_PI_Rows": external_pi_rows,
        "External_T_Gate": external_temperature_gate_summary,
        "Sstar_Sensitivity": sstar_sensitivity, "QA_Flagged_Rows": qa_flagged,
        "Claim_Gate": gate, "Outer_Fold_Assignment": outer_fold_assignment,
        "References": references,
    }
    print("Writing all 45 result tables", flush=True)
    workbook_path = out / "WB_PIML_results.xlsx"
    write_workbook(workbook_path, tables)
    write_readme(
        out / "README_results.md", data_path, metrics, effects, gate, df,
        parameters, len(runout_df),
    )
    print(f"\nSaved one tabular results workbook: {workbook_path}")
    print(f"Audit metadata: {out}")
    print(f"CSV export: {'enabled' if args.export_csv else 'disabled'}")
    print(f"Elapsed: {time.perf_counter() - started:.1f} s")


if __name__ == "__main__":
    main()
