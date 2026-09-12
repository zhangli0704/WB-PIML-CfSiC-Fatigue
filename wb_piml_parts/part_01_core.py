"""WB-PIML fatigue-life analysis with right-censored observations.

Walker--Basquin mechanical and environmental damage form the physical trunk.
Campaign-disjoint development folds select the residual family, gain, Walker
gamma and environmental regularization. Soft temperature support uses 250 C.
WB-Residual and Basquin-Residual provide independently selected simple-trunk controls.
Outputs include predictions, uncertainty intervals, validation, and sensitivity
results in a workbook, figures, and a presentation.
"""

from __future__ import annotations
from scipy.linalg import cho_factor, cho_solve
from sklearn.gaussian_process.kernels import Matern
import itertools

import argparse
import pickle
from collections import OrderedDict
from functools import wraps
import hashlib
import json
import math
import os
import platform
import tempfile
import time
import warnings
import zipfile
from dataclasses import dataclass, replace
from io import BytesIO
from pathlib import Path

os.environ.setdefault("MPLCONFIGDIR", os.path.join(os.environ.get("TEMP", "/tmp"), "piml_mpl"))

import matplotlib.pyplot as plt
from matplotlib.transforms import Bbox
import numpy as np
import pandas as pd
import scipy
import sklearn
from scipy.optimize import minimize
from joblib import Parallel, delayed, parallel_config
from threadpoolctl import threadpool_limits
import scipy.special as scipy_special
from sklearn.base import clone
from sklearn.compose import ColumnTransformer
from sklearn.ensemble import ExtraTreesRegressor, GradientBoostingRegressor, RandomForestRegressor
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import OneHotEncoder, StandardScaler
from sklearn.svm import SVR

try:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    Alignment = Font = PatternFill = get_column_letter = None

try:
    from PIL import Image
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt
except ImportError:
    Image = Presentation = RGBColor = PP_ALIGN = Inches = Pt = None


WORKERS = 8
TREE_JOBS = 1
BLAS_THREADS = 1

SEED = 20260824
GAMMA = 0.85
DAMAGE_CAP = 0.40
ARRHENIUS_REFERENCE_C = 800.0
ARRHENIUS_SCALE = 1000.0
LOG5 = math.log10(5.0)
SHEET = "Plot_Data_Verified"
EXPECTED_FILE = "data.xlsx"
EXPECTED_VERIFIED_ROWS = 223
EXPECTED_EXACT_ROWS = 158
EXPECTED_RUNOUT_ROWS = 65
DEFAULT_OUTPUT_DIR = Path(__file__).resolve().parent / "WB_PIML_results"
PPTX_OUTPUT_NAME = "WB_PIML_figures.pptx"
PRIMARY_GROUP_COLUMN = "campaign_id"
PROXIES = ["Dm", "Di", "Df", "Dox"]
TRAD_NUMERIC = ["logS", "R", "Tn", "logf", "logUTS", "Sa", "Smean"]
TRAD_CATEGORICAL = ["architecture", "architecture_detail", "env_class"]
NESTED_CANDIDATES = (
    # gamma, min_samples_leaf, max_features, residual gain eta
    (0.75, 2, 0.65, 0.50),
    (0.75, 3, 0.85, 1.00),
    (0.75, 5, 1.00, 0.50),
    (0.85, 2, 0.85, 1.00),
    (0.85, 3, 0.65, 0.50),
    (0.85, 3, 0.85, 0.00),
    (0.85, 3, 0.85, 0.50),
    (0.85, 3, 0.85, 1.00),
    (0.85, 5, 0.85, 1.00),
    (0.95, 2, 0.85, 1.00),
    (0.95, 3, 0.85, 0.50),
    (0.95, 3, 1.00, 1.00),
)
STRUCTURAL_GAMMAS = (.75, .85, .95)
ENVIRONMENT_STRENGTH_GRID = (0., .1, 1.)
RESIDUAL_GAIN_GRID = (.75, 1.)
HYBRID_FAMILIES = ("ET_residual", "physics_feature_ET_residual", "Matern_residual")
TEMPERATURE_POLICIES = ("soft250",)
def structural_candidate_grid():
    hybrids = [dict(pool="hybrid", family=family, gamma=gamma, strength=strength, eta=eta)
               for family, gamma, strength, eta in itertools.product(HYBRID_FAMILIES, STRUCTURAL_GAMMAS, ENVIRONMENT_STRENGTH_GRID, RESIDUAL_GAIN_GRID)]
    controls = [dict(pool="control", family="ET", leaf=leaf, max_features=mf)
                for leaf, mf in itertools.product((2, 3, 5, 8, 12, 16), (.65, .85, 1.))]
    controls += [dict(pool="control", family="GBDT", depth=depth, leaf=leaf, learning_rate=lr)
                 for depth, leaf, lr in itertools.product((1, 2, 3), (3, 8, 12), (.03, .07))]
    controls += [dict(pool="control", family="Matern", length_scale=length, alpha=alpha)
                 for length, alpha in itertools.product((1., 2., 4.), (.1, .3, 1., 3., 10., 30.))]
    controls += [dict(pool="control", family="fixed_ML_Ens")]
    controls += [dict(pool="control", family="physics_feature_ET", gamma=gamma, strength=strength)
                 for gamma, strength in itertools.product(STRUCTURAL_GAMMAS, ENVIRONMENT_STRENGTH_GRID)]
    return [dict(candidate_id=i, **config) for i, config in enumerate(hybrids + controls)]


STRUCTURAL_CANDIDATES = tuple(structural_candidate_grid())
HYBRID_CANDIDATES = tuple(c for c in STRUCTURAL_CANDIDATES if c["pool"] == "hybrid")
CONTROL_CANDIDATES = tuple(c for c in STRUCTURAL_CANDIDATES if c["pool"] == "control")
HYBRID_BASES = tuple({k:v for k,v in c.items() if k not in ("eta", "candidate_id")} for c in HYBRID_CANDIDATES[::2])

PHYSICS_MIX_GRID = (1.00,)
CALIBRATION_MODES = ("none",)
RMSE_TIE_TOLERANCE = 0.001
NEAR_TIE_RMSE = 0.020
# Match development exact-fracture temperatures within floating-point tolerance.
RESIDUAL_TEMPERATURE_MATCH_ATOL_C = 1e-9
CENSORED_MODELS = ("Lognormal-AFT", "Weibull-AFT")
REPEATED_VALIDATION_SEEDS = tuple(SEED + 101 * index for index in range(10))

# Source data for duplicate and label sensitivity comparisons.
AUDIT_CONTEXT: dict[str, object] = {}
ACTIVE_DATA_PATH: Path | None = None


@dataclass(frozen=True)
class PhysicalSpec:
    key: str
    channels: tuple[str, ...]
    mode: str = "linear"


@dataclass(frozen=True)
class CompetingDamageConfig:
    """Predeclared weak-physics assumptions used by the competing-damage fit.

    The environmental life-decrement penalty is selected in inner folds.
    Other bounds and ridge settings are fixed in the primary analysis;
    their alternatives are confined to sensitivity audits.
    """

    intercept_half_width: float = 0.75
    walker_half_width: float = 3.0
    walker_min: float = 0.05
    walker_max: float = 20.0
    environment_ratio_min: float = -8.0
    environment_ratio_max: float = 4.0
    environment_stress_min: float = 0.0
    environment_stress_max: float = 10.0
    arrhenius_min: float = 0.0
    arrhenius_max: float = 10.0
    ridge_strength: float = 0.01
    environment_strength: float = 0.0


DEFAULT_COMPETING_CONFIG = CompetingDamageConfig()


CLASSIC_WB_SPEC = PhysicalSpec("WB", ())
COMPETING_WB_SPEC = PhysicalSpec("WB-CD", (), "competing_damage")
PHYSICAL_SPECS = (
    CLASSIC_WB_SPEC,
    COMPETING_WB_SPEC,
    PhysicalSpec("WB-M", ("M",)),
    PhysicalSpec("WB-Ox", ("Ox",)),
)
EXTERNAL_MODELS = ("RF", "ExtraTrees", "ET-Walker", "GBDT", "SVR", "ML-Ens", "ML-Strong")
PRIMARY_COMPARATORS = ("ML-Ens", "ExtraTrees", "ET-Walker", "ML-Strong")
HYBRID_SPECS = (
    ("WB-PIML-Anchor", COMPETING_WB_SPEC),
    ("WB-PIML-M", PhysicalSpec("WB-M", ("M",))),
    ("WB-PIML-Ox", PhysicalSpec("WB-Ox", ("Ox",))),
)
PROPOSED_MODEL = "WB-PIML"
SIMPLE_CONTROL_MODELS = ("WB-Residual", "Basquin-Residual")


def legacy_first_models(models):
    """Append new controls without shifting existing bootstrap random streams."""
    names = set(models)
    return sorted(names - set(SIMPLE_CONTROL_MODELS)) + [
        name for name in SIMPLE_CONTROL_MODELS if name in names
    ]


def parse_bool(series: pd.Series, name: str = "boolean field") -> pd.Series:
    if pd.api.types.is_bool_dtype(series):
        return series.astype(bool)
    mapped = series.astype(str).str.strip().str.lower().map(
        {"true": True, "1": True, "yes": True, "y": True,
         "false": False, "0": False, "no": False, "n": False}
    )
    if mapped.isna().any():
        raise ValueError(f"{name} contains unrecognized values")
    return mapped.astype(bool)


def resolve_data(argument: str | None) -> Path:
    here = Path(__file__).resolve().parent
    candidates = [] if argument is None else [Path(argument)]
    candidates += [
        here / EXPECTED_FILE,
        Path.cwd() / EXPECTED_FILE,
        here / "upload" / EXPECTED_FILE,
    ]
    for path in candidates:
        if path.exists():
            return path.resolve()
    raise FileNotFoundError(f"Cannot find {EXPECTED_FILE}; pass --data")




def require_exact_training(frame: pd.DataFrame, context: str) -> pd.DataFrame:
    """Validate complementary exact/runout labels and unit likelihood weights."""
    if frame.empty:
        raise ValueError(f"{context}: empty training frame")
    if "is_runout" not in frame or "is_exact" not in frame:
        raise KeyError(f"{context}: missing exact/runout labels")
    if not frame["is_exact"].eq(~frame["is_runout"]).all():
        raise AssertionError(f"{context}: exact/runout labels are not complementary")
    if not frame["is_exact"].any():
        raise AssertionError(f"{context}: at least one exact failure is required")
    if "sample_weight" in frame.columns:
        supplied = pd.to_numeric(frame["sample_weight"], errors="coerce").to_numpy(float)
        if not np.isfinite(supplied).all() or not np.allclose(supplied, 1.0):
            raise ValueError(f"{context}: every analysis record must have weight 1")
    return frame


def sigmoid(x: np.ndarray) -> np.ndarray:
    return 1.0 / (1.0 + np.exp(-np.clip(x, -40.0, 40.0)))


def apply_physics_drivers(df: pd.DataFrame, water_factor: float = 1.5,
                          arrhenius_reference_c: float = ARRHENIUS_REFERENCE_C,
                          arrhenius_scale: float = ARRHENIUS_SCALE) -> pd.DataFrame:
    """Return a copy with explicit environmental/Arrhenius physics drivers.

    These constants are covariate-only assumptions.  Exposing them here makes
    the primary definition and every sensitivity variant auditable without
    changing labels, weights, folds, or validation responses.
    """
    if water_factor <= 0 or arrhenius_scale <= 0:
        raise ValueError("water_factor and arrhenius_scale must be positive")
    result = df.copy()
    env = result["env_class"].astype(str).str.lower()
    inert = env.str.contains("vacuum|inert|argon", regex=True, na=False)
    water = env.str.contains("water|steam", regex=True, na=False)
    result["environment_exposure"] = np.where(
        inert.to_numpy(bool), 0.0,
        np.where(water.to_numpy(bool), float(water_factor), 1.0),
    )
    temperature_k = result["T_C"].to_numpy(float) + 273.15
    reference_k = float(arrhenius_reference_c) + 273.15
    if reference_k <= 0 or np.any(temperature_k <= 0):
        raise ValueError("Arrhenius temperatures must be above absolute zero")
    result["arrhenius_temperature_drive"] = float(arrhenius_scale) * (
        1.0 / reference_k - 1.0 / temperature_k
    )
    return result


def load_data(path: Path) -> pd.DataFrame:
    if SHEET not in pd.ExcelFile(path).sheet_names:
        raise KeyError(f"Missing sheet {SHEET!r}")
    configure_audit_context(path)
    df = pd.read_excel(path, sheet_name=SHEET).copy()
    required = {
        "row_id", "source_id", "architecture", "arch_group", "env_class", "R", "T_C",
        "frequency_Hz", "Nf_cycles", "UTS_MPa", "stress_level_sigma_max_over_UTS",
        "sigma_max_MPa_before_normalization", "is_runout", "is_exact_failure",
        "record_equivalence_id", "campaign_id", "training_role", "sample_weight",
        "event_observed", "censoring_type", "likelihood_weight", "point_metric_eligible",
    }
    missing = sorted(required - set(df.columns))
    if missing:
        raise KeyError(f"Missing required columns: {missing}")
    df["row_id"] = pd.to_numeric(df["row_id"], errors="raise").astype(int)
    if df["row_id"].duplicated().any():
        raise ValueError("row_id must be unique")
    df["source_id"] = df["source_id"].fillna("").astype(str).str.strip()
    if df["source_id"].eq("").any():
        raise ValueError("source_id contains blanks")
    df["architecture_detail"] = df["architecture"].fillna("unknown").astype(str).str.strip().str.lower()
    df["architecture"] = (df["arch_group"].fillna(df["architecture_detail"]).fillna("unknown")
                          .astype(str).str.strip().str.lower())
    df["env_class"] = df["env_class"].fillna("unknown").astype(str).str.strip().str.lower()
    if "temp_bin" in df.columns:
        df["temp_bin"] = df["temp_bin"].fillna("unknown").astype(str).str.strip().str.lower()
    else:
        # Use exact temperatures when no temperature-group labels are supplied.
        df["temp_bin"] = pd.to_numeric(df["T_C"], errors="coerce").map(lambda value: f"{value:g}c")
    for col in ["R", "T_C", "frequency_Hz", "Nf_cycles", "UTS_MPa",
                "stress_level_sigma_max_over_UTS", "sigma_max_MPa_before_normalization"]:
        df[col] = pd.to_numeric(df[col], errors="coerce")
    numeric = ["R", "T_C", "frequency_Hz", "Nf_cycles", "UTS_MPa",
               "stress_level_sigma_max_over_UTS", "sigma_max_MPa_before_normalization"]
    if df[numeric].isna().any().any() or not np.isfinite(df[numeric].to_numpy(float)).all():
        raise ValueError("Required numeric columns contain missing/non-numeric values")
    # The 800 C boundary belongs to the 800-1000 C regime for every source.
    boundary = df["T_C"].eq(800) & df["temp_bin"].isin(["500-800", "800-1000"])
    df.loc[boundary, "temp_bin"] = "800-1000"
    temperature_groups = df.groupby("T_C")["temp_bin"].nunique()
    if temperature_groups.gt(1).any():
        raise ValueError("Each temperature must belong to one temp_bin; conflicting temperatures: "
                         f"{temperature_groups.index[temperature_groups.gt(1)].tolist()}")
    if (df[["frequency_Hz", "Nf_cycles", "UTS_MPa", "stress_level_sigma_max_over_UTS"]] <= 0).any().any():
        raise ValueError("Frequency, life, UTS and normalized stress must be positive")
    if (df["R"] >= 1.0).any():
        raise ValueError("Walker correction requires every stress ratio R < 1")
    df["is_runout"] = parse_bool(df["is_runout"], "is_runout")
    supplied_exact = parse_bool(df["is_exact_failure"], "is_exact_failure")
    mismatch = supplied_exact.ne(~df["is_runout"])
    if mismatch.any():
        bad_rows = df.loc[mismatch, "row_id"].astype(int).tolist()
        raise ValueError(
            "is_exact_failure must equal ~is_runout; mismatched row_id values: "
            f"{bad_rows[:20]}"
        )
    df["supplied_is_runout"] = df["is_runout"].astype(bool)
    df["supplied_is_exact_failure"] = supplied_exact.astype(bool)
    df["label_override_applied"] = False
    df["label_review_status"] = "accepted_from_Plot_Data_Verified"
    if "evidence_source" in df.columns:
        df["label_review_evidence"] = df["evidence_source"].fillna("").astype(str)
    else:
        df["label_review_evidence"] = "verified workbook label retained without code-side override"
    df["reviewed_is_exact_failure"] = ~df["is_runout"]
    df["is_exact"] = ~df["is_runout"]
    df["label_consistent"] = df["reviewed_is_exact_failure"].eq(~df["is_runout"])
    df["training_role"] = df["training_role"].fillna("").astype(str).str.strip()
    df["sample_weight"] = pd.to_numeric(df["sample_weight"], errors="coerce")
    expected_role = np.where(
        df["is_exact"], "exact_likelihood_training", "right_censored_likelihood_training"
    )
    role_mismatch = df["training_role"].ne(expected_role)
    if role_mismatch.any():
        bad = df.loc[role_mismatch, ["row_id", "training_role", "is_runout"]]
        raise ValueError(
            "training_role is inconsistent with the joint-likelihood contract: "
            f"{bad.to_dict('records')[:20]}"
        )
    df["likelihood_weight"] = pd.to_numeric(df["likelihood_weight"], errors="coerce")
    if not np.allclose(df["sample_weight"], 1.0) or not np.allclose(df["likelihood_weight"], 1.0):
        bad = df.loc[
            ~np.isclose(df["sample_weight"], 1.0) | ~np.isclose(df["likelihood_weight"], 1.0),
            ["row_id", "training_role", "sample_weight", "likelihood_weight"],
        ]
        raise ValueError(f"Every unique record must have unit likelihood weight: {bad.to_dict('records')[:20]}")
    event = pd.to_numeric(df["event_observed"], errors="coerce")
    point = pd.to_numeric(df["point_metric_eligible"], errors="coerce")
    expected_event = df["is_exact"].astype(int)
    expected_censoring = np.where(df["is_exact"], "exact", "right_censored")
    if not event.eq(expected_event).all() or not point.eq(expected_event).all():
        raise ValueError("event_observed/point_metric_eligible are inconsistent with exact/runout labels")
    if not df["censoring_type"].astype(str).eq(expected_censoring).all():
        raise ValueError("censoring_type is inconsistent with exact/runout labels")
    if "include_in_verified_model" in df.columns:
        included = parse_bool(df["include_in_verified_model"], "include_in_verified_model")
        if not included.all():
            bad = df.loc[~included, "row_id"].astype(int).tolist()
            raise ValueError(
                "Plot_Data_Verified may contain only retained unique records; "
                f"excluded row_id values were found: {bad[:20]}"
            )
    df["logN"] = np.log10(df["Nf_cycles"].clip(lower=1.0))
    df["S"] = df["stress_level_sigma_max_over_UTS"]
    df["logS"] = np.log10(df["S"].clip(lower=1e-8))
    df["logf"] = np.log10(df["frequency_Hz"].clip(lower=1e-8))
    df["logUTS"] = np.log10(df["UTS_MPa"].clip(lower=1e-8))
    df["Tn"] = (df["T_C"] - 25.0) / 1000.0
    df["Sa"] = 0.5 * df["S"] * np.clip(1.0 - df["R"], 0.0, 2.0)
    df["Smean"] = 0.5 * df["S"] * (1.0 + df["R"])

    df = apply_physics_drivers(df)
    env = df["env_class"].str.lower()
    inert = env.str.contains("vacuum|inert|argon", regex=True, na=False)
    water = env.str.contains("water|steam", regex=True, na=False)
    thermal = sigmoid((df["T_C"].to_numpy(float) - 800.0) / 200.0)
    slow = sigmoid(-df["logf"].to_numpy(float))
    oxidation = (~inert).to_numpy(float) * thermal * (0.5 + 0.5 * slow)
    oxidation = np.maximum(oxidation, water.to_numpy(float) * thermal)
    df["oxidation"] = np.clip(oxidation, 0.0, 1.0)
    stress = np.clip(df["S"].to_numpy(float), 0.0, 1.0)
    df["Dm"] = stress
    df["Di"] = np.clip(stress * (0.60 + 0.40 * df["oxidation"]), 0.0, 1.0)
    df["Df"] = np.clip(stress**2, 0.0, 1.0)
    df["Dox"] = df["oxidation"]
    df["M"] = df[["Dm", "Di", "Df"]].mean(axis=1)
    df["Ox"] = df["Dox"]
    df["Q"] = df[PROXIES].mean(axis=1)
    series_columns = ["source_id", "architecture_detail", "env_class", "T_C", "R", "UTS_MPa"]
    df["series_group_id"] = df[series_columns].astype(str).agg("|".join, axis=1)
    if "raw_Nf" in df.columns:
        raw_nf = df["raw_Nf"].fillna("").astype(str).str.strip().str.lower()
    else:
        raw_nf = pd.Series("", index=df.index)
    round_limit = np.isclose(df["Nf_cycles"], 1e6) | np.isclose(df["Nf_cycles"], 1e7)
    df["qa_exact_round_limit"] = df["is_exact"] & round_limit & ~raw_nf.str.contains("~|未断|runout|>", regex=True)
    df["qa_sstar_gt_1"] = df["S"] > 1.0
    return add_campaign_metadata(df.reset_index(drop=True))


def source_weights(df: pd.DataFrame) -> np.ndarray:
    count = df.groupby("source_id")["source_id"].transform("size").to_numpy(float)
    weight = 1.0 / count
    return weight / np.mean(weight)


def training_weights(df: pd.DataFrame) -> np.ndarray:
    """Return strict unit likelihood weights for all unique records.

    Row numbers, publication sources, campaigns and equivalence identifiers
    never alter a training record's importance.  Exact rows enter via density
    and runouts via survival, but both have likelihood weight one.
    """
    require_exact_training(df, "training_weights")
    if "sample_weight" in df.columns:
        supplied = pd.to_numeric(df["sample_weight"], errors="coerce").to_numpy(float)
        if not np.isfinite(supplied).all() or not np.allclose(supplied, 1.0):
            bad = df.loc[
                ~np.isclose(supplied, 1.0), ["row_id", "sample_weight"]
            ].to_dict("records")
            raise ValueError(f"Every unique likelihood record must have weight 1: {bad[:20]}")
    return np.ones(len(df), dtype=float)


def joint_aft_nll(frame: pd.DataFrame, mu: np.ndarray, sigma: float,
                  weights: np.ndarray | None = None) -> dict[str, float]:
    """Lognormal AFT NLL with exact density and right-censored survival terms."""
    require_exact_training(frame, "joint_aft_nll")
    if not np.isfinite(sigma) or sigma <= 0:
        raise ValueError("AFT sigma must be positive and finite")
    location = np.asarray(mu, float)
    if len(location) != len(frame) or not np.isfinite(location).all():
        raise ValueError("AFT location prediction is missing or non-finite")
    weight = training_weights(frame) if weights is None else np.asarray(weights, float)
    y = frame["logN"].to_numpy(float)
    exact_mask = frame["is_exact"].to_numpy(bool)
    runout_mask = ~exact_mask
    z = (y - location) / sigma
    losses = np.empty(len(frame), float)
    losses[exact_mask] = (
        0.5 * math.log(2.0 * math.pi) + math.log(sigma)
        + 0.5 * z[exact_mask] ** 2
    )
    # S(t|mu,sigma) = Phi((mu-t)/sigma); log_ndtr is stable in the tail.
    losses[runout_mask] = -scipy_special.log_ndtr(-z[runout_mask])
    total = float(np.sum(weight * losses) / np.sum(weight))
    exact_nll = float(np.average(losses[exact_mask], weights=weight[exact_mask]))
    runout_nll = (
        float(np.average(losses[runout_mask], weights=weight[runout_mask]))
        if runout_mask.any() else np.nan
    )
    return {
        "joint_censored_NLL": total,
        "exact_density_NLL": exact_nll,
        "runout_survival_NLL": runout_nll,
    }


def analysis_weights(df: pd.DataFrame) -> np.ndarray:
    """Campaign-balanced weights for reported metrics and diagnostics only.

    These weights never enter model fitting.  ``Plot_Data_Verified`` is already
    physically deduplicated, so each equivalence ID occurs only once.
    """
    required = {"campaign_id", "record_equivalence_id"}
    if not required.issubset(df.columns):
        return source_weights(df)
    frame = df[["campaign_id", "record_equivalence_id"]].astype(str)
    pair_size = frame.groupby(
        ["campaign_id", "record_equivalence_id"]
    )["record_equivalence_id"].transform("size").to_numpy(float)
    equivalence_per_campaign = (
        frame.drop_duplicates()
        .groupby("campaign_id")["record_equivalence_id"].nunique()
    )
    n_equivalence = frame["campaign_id"].map(equivalence_per_campaign).to_numpy(float)
    weight = 1.0 / (pair_size * n_equivalence)
    return weight / np.mean(weight)


def grouped_folds(df: pd.DataFrame, n_splits: int, seed: int):
    """Primary campaign-disjoint folds used by every nested model step."""
    require_exact_training(df, "grouped_folds")
    group_column = PRIMARY_GROUP_COLUMN
    summary = df.groupby(group_column).agg(n=("row_id", "size")).reset_index()
    n_splits = min(n_splits, len(summary))
    rng = np.random.default_rng(seed)
    summary["tie"] = rng.random(len(summary))
    summary = summary.sort_values(["n", "tie"], ascending=[False, True])
    target = float(summary["n"].sum()) / n_splits
    load = np.zeros(n_splits)
    assignment: dict[str, int] = {}
    for _, row in summary.iterrows():
        score = load / max(target, 1.0)
        fold = int(np.argmin(score))
        load[fold] += float(row["n"])
        assignment[str(row[group_column])] = fold
    fold_id = df[group_column].astype(str).map(assignment).to_numpy(int)
    splits = []
    for fold in range(n_splits):
        valid = np.flatnonzero(fold_id == fold)
        train = np.flatnonzero(fold_id != fold)
        if not len(train) or not len(valid):
            raise RuntimeError("Empty grouped fold")
        if set(df.iloc[train][group_column]) & set(df.iloc[valid][group_column]):
            raise AssertionError("Campaign leakage")
        if not df.iloc[train]["is_exact"].any() or not df.iloc[valid]["is_exact"].any():
            raise RuntimeError("Every outer fold must contain exact failures")
        splits.append((train, valid))
    return splits, fold_id + 1


def walker_log_stress(df: pd.DataFrame, gamma: float = GAMMA) -> np.ndarray:
    factor = np.clip(0.5 * (1.0 - df["R"].to_numpy(float)), 1e-6, None)
    return df["logS"].to_numpy(float) + gamma * np.log10(factor)


def physical_design(df: pd.DataFrame, spec: PhysicalSpec, gamma: float = GAMMA) -> np.ndarray:
    if spec.mode != "linear":
        raise ValueError(f"{spec.key} uses a nonlinear competing-damage equation")
    columns = [np.ones(len(df)), -walker_log_stress(df, gamma)]
    columns += [-df[channel].to_numpy(float) for channel in spec.channels]
    return np.column_stack(columns)


def _fit_linear_physical(frame: pd.DataFrame, spec: PhysicalSpec,
                         gamma: float) -> np.ndarray:
    """Fit a physical location with joint exact/right-censored AFT likelihood."""
    require_exact_training(frame, f"_fit_linear_physical[{spec.key}]")
    x = physical_design(frame, spec, gamma)
    exact_mask = frame["is_exact"].to_numpy(bool)
    initial_coef = np.linalg.lstsq(
        x[exact_mask], frame.loc[exact_mask, "logN"].to_numpy(float), rcond=None
    )[0]
    initial_coef[1] = np.clip(initial_coef[1], 0.05, 20.0)
    if len(initial_coef) > 2:
        initial_coef[2:] = np.clip(initial_coef[2:], 0.0, DAMAGE_CAP)
    initial_sigma = np.clip(
        np.std(frame.loc[exact_mask, "logN"].to_numpy(float) - x[exact_mask] @ initial_coef),
        0.10, 3.0,
    )

    def loss(parameter: np.ndarray) -> float:
        coef = parameter[:-1]
        sigma = float(np.exp(parameter[-1]))
        nll = joint_aft_nll(frame, x @ coef, sigma)["joint_censored_NLL"]
        ridge = 1e-4 * float(np.sum(coef[2:] ** 2)) if len(coef) > 2 else 0.0
        return nll + ridge

    initial = np.r_[initial_coef, math.log(initial_sigma)]
    bounds = ([(None, None), (0.05, 20.0)]
              + [(0.0, DAMAGE_CAP)] * len(spec.channels)
              + [(math.log(0.08), math.log(4.0))])
    result = minimize(
        loss, initial, method="L-BFGS-B", bounds=bounds,
        options={"maxiter": 1200, "ftol": 1e-12, "gtol": 1e-8},
    )
    if not result.success or not np.isfinite(result.fun):
        # Retry the same objective and bounds with Powell after L-BFGS-B failure.
        retry_start = (
            np.asarray(result.x, float)
            if np.size(result.x) == np.size(initial)
            and np.all(np.isfinite(result.x))
            else initial
        )
        fallback = minimize(
            loss, retry_start, method="Powell", bounds=bounds,
            options={"maxiter": 6000, "xtol": 1e-9, "ftol": 1e-10},
        )
        if fallback.success and np.isfinite(fallback.fun):
            result = fallback
        else:
            raise RuntimeError(
                f"{spec.key} censored-likelihood optimization failed: "
                f"L-BFGS-B={result.message}; Powell={fallback.message}"
            )
    return np.asarray(result.x[:-1], float)


def _predict_competing_damage(df: pd.DataFrame, coef: np.ndarray,
                              gamma: float) -> np.ndarray:
    """Predict log10 life from competing per-cycle damage rates.

    ``coef = [a, b, log10_K_env, m, q_T]``.  The two damage
    channels are strictly non-negative.  Positive ``b`` and ``m`` enforce a
    non-increasing stress-life relation, positive ``q_T`` increases the
    environmental damage rate with temperature, and ``1/f`` implements
    exposure time per cycle.
    """
    if len(coef) != 5:
        raise ValueError("The competing-damage trunk requires five parameters")
    intercept, b, log_k, m, q_t = np.asarray(coef, float)
    log_sw = walker_log_stress(df, gamma)
    log_d_mech = -intercept + b * log_sw

    exposure = df["environment_exposure"].to_numpy(float)
    active = exposure > 0.0
    log_d_env = np.full(len(df), -np.inf, dtype=float)
    log_d_env[active] = (
        log_d_mech[active]
        + log_k
        + np.log10(exposure[active])
        + (q_t * df.loc[active, "arrhenius_temperature_drive"].to_numpy(float)
           / math.log(10.0))
        + m * log_sw[active]
        - df.loc[active, "logf"].to_numpy(float)
    )
    # Stable base-10 log-sum-exp: log10(10**x + 10**y).
    maximum = np.maximum(log_d_mech, log_d_env)
    log_total_damage = maximum + np.log10(
        np.power(10.0, log_d_mech - maximum)
        + np.power(10.0, log_d_env - maximum)
    )
    return -log_total_damage


def cache_physical_fit(function, validate, damage_cap, maxsize=2048):
    cache = OrderedDict()
    stats = {'hits': 0, 'misses': 0}
    gamma_default, config_default = function.__defaults__

    def array_token(values):
        array = np.asarray(values)
        content = (
            tuple(pickle.dumps(value, protocol=5) for value in array.flat)
            if array.dtype.hasobject else array.tobytes(order='C')
        )
        return array.dtype.str, array.shape, content

    def key_for(train, spec, gamma, competing_config):
        dtypes = tuple(map(str, train.dtypes))
        payload = (
            tuple(train.columns.names), array_token(train.columns),
            tuple(train.index.names), array_token(train.index),
            tuple((dtypes[i], array_token(train.iloc[:, i].to_numpy()))
                  for i in range(train.shape[1])),
            (spec.key, tuple(spec.channels), spec.mode), gamma,
            tuple(sorted(vars(competing_config).items())), damage_cap(),
        )
        return hashlib.sha256(pickle.dumps(payload, protocol=5)).digest()

    @wraps(function)
    def cached(train, spec, gamma=gamma_default, competing_config=config_default):
        validate(train, f"fit_physical[{spec.key}]")
        key = key_for(train, spec, gamma, competing_config)
        if key in cache:
            stats['hits'] += 1
            cache.move_to_end(key)
            return cache[key].copy()
        result = function(train, spec, gamma, competing_config)
        stats['misses'] += 1
        cache[key] = result.copy()
        if len(cache) > maxsize:
            cache.popitem(last=False)
        return result

    cached.cache_key = key_for
    cached.cache_info = lambda: dict(**stats, size=len(cache))
    cached.cache_clear = cache.clear
    return cached


def _fit_physical_uncached(train: pd.DataFrame, spec: PhysicalSpec, gamma: float = GAMMA,
                 competing_config: CompetingDamageConfig = DEFAULT_COMPETING_CONFIG,
                 ) -> np.ndarray:
    require_exact_training(train, f"fit_physical[{spec.key}]")
    exact = train.loc[train["is_exact"]].reset_index(drop=True)
    if spec.mode == "linear":
        return _fit_linear_physical(train, spec, gamma)
    if spec.mode != "competing_damage":
        raise ValueError(f"Unknown physical mode: {spec.mode}")

    classic = _fit_linear_physical(train, CLASSIC_WB_SPEC, gamma)
    classic_intercept = float(classic[0])
    classic_b = float(classic[1])

    def loss(parameter: np.ndarray) -> float:
        coef = parameter[:-1]
        sigma = float(np.exp(parameter[-1]))
        prediction = _predict_competing_damage(train, coef, gamma)
        nll = joint_aft_nll(train, prediction, sigma)["joint_censored_NLL"]
        # Regularize the mechanical limit toward the independently fitted WB model.
        ridge = float(competing_config.ridge_strength) * (
            (coef[0] - classic_intercept) ** 2
            + 0.10 * (coef[1] - classic_b) ** 2
            + 0.01 * coef[3] ** 2
            + 0.01 * coef[4] ** 2
        )
        if competing_config.environment_strength == 0.0:
            return nll + ridge
        decrement = environmental_decrement(train, coef, gamma)
        return nll + ridge + competing_config.environment_strength * float(np.mean(np.square(decrement)))

    bounds = [
        (classic_intercept - competing_config.intercept_half_width,
         classic_intercept + competing_config.intercept_half_width),
        (max(competing_config.walker_min,
             classic_b - competing_config.walker_half_width),
         min(competing_config.walker_max,
             classic_b + competing_config.walker_half_width)),
        (competing_config.environment_ratio_min,
         competing_config.environment_ratio_max),
        (competing_config.environment_stress_min,
         competing_config.environment_stress_max),
        (competing_config.arrhenius_min,
         competing_config.arrhenius_max),
    ]
    starts = [
        np.array([classic_intercept, classic_b, -2.0, 1.0, 4.0, math.log(0.8)]),
        np.array([classic_intercept, classic_b, -4.0, 2.0, 7.0, math.log(1.2)]),
        np.array([classic_intercept, classic_b, 0.0, 4.0, 2.0, math.log(1.0)]),
    ]
    bounds = bounds + [(math.log(0.08), math.log(4.0))]
    results = [
        minimize(loss, np.clip(start, [b[0] for b in bounds], [b[1] for b in bounds]),
                 method="L-BFGS-B", bounds=bounds,
                 options={"maxiter": 1200, "ftol": 1e-12, "gtol": 1e-8})
        for start in starts
    ]
    valid = [result for result in results if result.success and np.isfinite(result.fun)]
    if not valid:
        messages = "; ".join(str(result.message) for result in results)
        raise RuntimeError(f"{spec.key} optimization failed: {messages}")
    best = min(valid, key=lambda result: float(result.fun))
    return np.asarray(best.x[:-1], float)


fit_physical = cache_physical_fit(
    _fit_physical_uncached, require_exact_training, lambda: DAMAGE_CAP
)


def predict_physical(df: pd.DataFrame, spec: PhysicalSpec, coef: np.ndarray,
                     gamma: float = GAMMA) -> np.ndarray:
    if spec.mode == "competing_damage":
        return _predict_competing_damage(df, coef, gamma)
    return physical_design(df, spec, gamma) @ coef


def physical_parameter_values(spec: PhysicalSpec, coef: np.ndarray) -> dict[str, float]:
    """Return named parameters without pretending nonlinear coefficients are intercepts."""
    values = {
        "intercept": np.nan,
        "walker_slope": np.nan,
        "mechanical_log10_damage_scale": np.nan,
        "environment_log10_damage_ratio": np.nan,
        "environment_stress_exponent": np.nan,
        "arrhenius_temperature_sensitivity": np.nan,
        "M_coefficient": np.nan,
        "Ox_coefficient": np.nan,
        "Q_coefficient": np.nan,
    }
    if spec.mode == "competing_damage":
        values.update({
            "intercept": float(coef[0]),
            "mechanical_log10_damage_scale": float(-coef[0]),
            "walker_slope": float(coef[1]),
            "environment_log10_damage_ratio": float(coef[2]),
            "environment_stress_exponent": float(coef[3]),
            "arrhenius_temperature_sensitivity": float(coef[4]),
        })
    else:
        values["intercept"] = float(coef[0])
        values["walker_slope"] = float(coef[1])
        for channel, value in zip(spec.channels, coef[2:]):
            values[f"{channel}_coefficient"] = float(value)
    return values


def inner_oof(train: pd.DataFrame, spec: PhysicalSpec, seed: int) -> pd.DataFrame:
    require_exact_training(train, f"inner_oof[{spec.key}]")
    splits, _ = grouped_folds(train, min(4, train[PRIMARY_GROUP_COLUMN].nunique()), seed)
    pieces = []
    for fold, (fit_idx, valid_idx) in enumerate(splits, 1):
        fit = train.iloc[fit_idx].reset_index(drop=True)
        valid = train.iloc[valid_idx].reset_index(drop=True)
        coef = fit_physical(fit, spec)
        part = valid[["row_id", "source_id", "campaign_id", "record_equivalence_id",
                      "logN", "is_exact", "is_runout"]].copy()
        part["inner_fold"] = fold
        part["mu"] = predict_physical(valid, spec, coef)
        pieces.append(part)
    oof = pd.concat(pieces, ignore_index=True)
    if oof["row_id"].duplicated().any() or set(oof["row_id"]) != set(train["row_id"]):
        raise AssertionError("Inner OOF must predict each development row once")
    return oof


def calibrate_aft_scale(oof: pd.DataFrame) -> dict[str, float]:
    require_exact_training(oof, "calibrate_aft_scale")
    exact = oof.loc[oof["is_exact"]].reset_index(drop=True)
    ew = training_weights(exact)

    def objective(value: np.ndarray) -> float:
        sigma = float(np.exp(value[0]))
        return joint_aft_nll(oof, oof["mu"].to_numpy(float), sigma)["joint_censored_NLL"]

    start = math.log(np.clip(np.sqrt(np.average((exact["logN"] - exact["mu"]) ** 2, weights=ew)), 0.10, 3.0))
    result = minimize(objective, np.array([start]), method="L-BFGS-B", bounds=[(math.log(0.08), math.log(4.0))])
    if not result.success:
        raise RuntimeError("AFT scale calibration failed")
    sigma = float(np.exp(result.x[0]))
    losses = joint_aft_nll(oof, oof["mu"].to_numpy(float), sigma)
    return {"sigma": sigma, "calibration_objective": float(result.fun), **losses}


def weighted_quantile(values: np.ndarray, quantile: float, weights: np.ndarray) -> float:
    order = np.argsort(values)
    values = np.asarray(values, float)[order]
    weights = np.asarray(weights, float)[order]
    cumulative = np.cumsum(weights) / np.sum(weights)
    return float(np.interp(quantile, cumulative, values))




def make_preprocessor(extra_numeric: tuple[str, ...] = ()) -> ColumnTransformer:
    try:
        encoder = OneHotEncoder(handle_unknown="ignore", sparse_output=False)
    except TypeError:  # scikit-learn < 1.2
        encoder = OneHotEncoder(handle_unknown="ignore", sparse=False)
    return ColumnTransformer(
        [("numeric", StandardScaler(), TRAD_NUMERIC + list(extra_numeric)),
         ("categorical", encoder, TRAD_CATEGORICAL)],
        sparse_threshold=0.0,
    )






def censored_hinge_rmse(frame: pd.DataFrame, prediction_column: str = "mu") -> float:
    """Campaign/equivalence-balanced one-sided error for right-censored rows."""
    runout = frame.loc[frame["is_runout"]].reset_index(drop=True)
    if not len(runout):
        return 0.0
    violation = np.maximum(
        runout["logN"].to_numpy(float) - runout[prediction_column].to_numpy(float),
        0.0,
    )
    weight = analysis_weights(runout)
    return float(np.sqrt(np.sum(weight * violation**2) / np.sum(weight)))


def choose_rmse_first(candidates: pd.DataFrame) -> pd.Series:
    """Protect exact RMSE, then prefer the lowest joint censored NLL."""
    best_rmse = float(candidates["inner_SB_RMSE"].min())
    eligible = candidates.loc[
        candidates["inner_SB_RMSE"] <= best_rmse + RMSE_TIE_TOLERANCE
    ].copy()
    sort_columns = [
        column for column in (
            "inner_joint_censored_NLL", "calibration_complexity",
            "physics_mix", "min_samples_leaf", "candidate_index",
        ) if column in eligible.columns
    ]
    if not sort_columns:
        sort_columns = ["inner_SB_RMSE"]
    return eligible.sort_values(sort_columns, na_position="last").iloc[0]


def apply_safe_fusion(et_prediction: np.ndarray, physical_prediction: np.ndarray,
                      physics_mix: float, calibration_intercept: float,
                      calibration_slope: float) -> np.ndarray:
    """Return the physical hybrid after validating unit weight and no calibration."""
    if not (
        np.isclose(physics_mix, 1.0)
        and np.isclose(calibration_intercept, 0.0)
        and np.isclose(calibration_slope, 1.0)
    ):
        raise AssertionError("WB-PIML cannot switch off physics or calibrate outer predictions")
    if len(et_prediction) != len(physical_prediction):
        raise AssertionError("Safe-fusion branches must cover identical rows")
    return np.asarray(physical_prediction, float)


def select_safe_fusion(physical_oof: pd.DataFrame, et_oof: pd.DataFrame
                       ) -> tuple[dict[str, float | str], pd.DataFrame, pd.DataFrame]:
    """Finalize the residual branch with unit physical weight and no calibration.

    ExtraTrees OOF rows are used only to check matching development-row coverage.
    """
    if PHYSICS_MIX_GRID != (1.0,) or CALIBRATION_MODES != ("none",):
        raise AssertionError("WB-PIML requires an always-active, uncalibrated physical trunk")
    metadata = [
        "row_id", "source_id", "campaign_id", "record_equivalence_id",
        "logN", "is_exact", "is_runout", "inner_fold",
    ]
    physical = physical_oof[metadata + ["mu"]].rename(columns={"mu": "physical_mu"})
    direct = et_oof[["row_id", "mu"]].rename(columns={"mu": "et_mu"})
    merged = physical.merge(direct, on="row_id", how="inner", validate="one_to_one")
    if len(merged) != len(physical_oof):
        raise AssertionError("Safe-fusion OOF branches do not cover identical development rows")

    candidate = merged.copy()
    candidate["mu"] = candidate["physical_mu"].to_numpy(float)
    exact = candidate.loc[candidate["is_exact"]].copy()
    exact["pred_logN"] = exact["mu"]
    exact_score = sb_metrics([
        group.reset_index(drop=True)
        for _, group in exact.groupby(PRIMARY_GROUP_COLUMN)
    ])["SB_RMSE"]
    censor_score = censored_hinge_rmse(candidate)
    scale = calibrate_aft_scale(candidate)
    candidates = pd.DataFrame([{
        "physics_mix": 1.0,
        "calibration_mode": "none",
        "calibration_intercept": 0.0,
        "calibration_slope": 1.0,
        "calibration_complexity": 0,
        "inner_SB_RMSE": float(exact_score),
        "inner_runout_hinge_RMSE": float(censor_score),
        "inner_joint_censored_NLL": float(scale["joint_censored_NLL"]),
        "within_RMSE_tolerance": True,
    }])
    best = candidates.iloc[0]
    selected_oof = candidate[metadata + ["mu"]].copy()
    choice: dict[str, float | str] = {
        "physics_mix": 1.0,
        "calibration_mode": "none",
        "calibration_intercept": float(best["calibration_intercept"]),
        "calibration_slope": float(best["calibration_slope"]),
        "inner_SB_RMSE": float(best["inner_SB_RMSE"]),
        "inner_runout_hinge_RMSE": float(best["inner_runout_hinge_RMSE"]),
        "inner_joint_censored_NLL": float(best["inner_joint_censored_NLL"]),
        "RMSE_tolerance": RMSE_TIE_TOLERANCE,
    }
    candidates = candidates.sort_values(
        ["within_RMSE_tolerance", "inner_SB_RMSE", "inner_joint_censored_NLL"],
        ascending=[False, True, True],
    ).reset_index(drop=True)
    return choice, selected_oof, candidates




def fit_extra_trees(train: pd.DataFrame, valid: pd.DataFrame, seed: int, trees: int,
                    min_samples_leaf: int = 3,
                    max_features: float = 0.85) -> np.ndarray:
    require_exact_training(train, "fit_extra_trees")
    exact = train.loc[train["is_exact"]].reset_index(drop=True)
    estimator = ExtraTreesRegressor(
        n_estimators=trees, min_samples_leaf=min_samples_leaf,
        max_features=max_features, random_state=seed, n_jobs=TREE_JOBS,
    )
    pipe = Pipeline([("preprocess", clone(make_preprocessor())), ("model", estimator)])
    pipe.fit(exact, exact["logN"], model__sample_weight=training_weights(exact))
    return np.asarray(pipe.predict(valid), float)


def nested_select_extra_trees(train: pd.DataFrame, seed: int, trees: int
                              ) -> tuple[dict[str, float], pd.DataFrame, pd.DataFrame]:
    require_exact_training(train, "nested_select_extra_trees")
    splits, _ = grouped_folds(train, min(4, train[PRIMARY_GROUP_COLUMN].nunique()), seed)
    unique_configs = sorted({(leaf, features) for _, leaf, features, _ in NESTED_CANDIDATES})
    rows = []
    stored: dict[int, pd.DataFrame] = {}
    for config_index, (min_leaf, max_features) in enumerate(unique_configs):
        pieces = []
        for inner_fold, (fit_idx, valid_idx) in enumerate(splits, 1):
            fit = train.iloc[fit_idx].reset_index(drop=True)
            valid = train.iloc[valid_idx].reset_index(drop=True)
            mu = fit_extra_trees(
                fit, valid, seed + 10000 * config_index + 100 * inner_fold + 7,
                min(trees, 300), min_leaf, max_features,
            )
            part = valid[["row_id", "source_id", "campaign_id", "record_equivalence_id",
                          "logN", "is_exact", "is_runout"]].copy()
            part["inner_fold"] = inner_fold
            part["mu"] = mu
            pieces.append(part)
        oof = pd.concat(pieces, ignore_index=True)
        exact = oof.loc[oof["is_exact"]].copy()
        exact["pred_logN"] = exact["mu"]
        score = sb_metrics([g.reset_index(drop=True) for _, g in exact.groupby(PRIMARY_GROUP_COLUMN)])["SB_RMSE"]
        rows.append({"config_index": config_index, "min_samples_leaf": min_leaf,
                     "max_features": max_features, "inner_SB_RMSE": score})
        stored[config_index] = oof
    candidates = pd.DataFrame(rows).sort_values(
        ["inner_SB_RMSE", "min_samples_leaf", "config_index"]
    ).reset_index(drop=True)
    best = candidates.iloc[0]
    chosen = {"config_index": int(best["config_index"]),
              "min_samples_leaf": int(best["min_samples_leaf"]),
              "max_features": float(best["max_features"]),
              "inner_SB_RMSE": float(best["inner_SB_RMSE"])}
    return chosen, stored[chosen["config_index"]], candidates


def fit_et_walker(train: pd.DataFrame, valid: pd.DataFrame, seed: int, trees: int,
                  gamma: float, min_samples_leaf: int,
                  max_features: float) -> np.ndarray:
    """Exact-only direct ExtraTrees control with a Walker-stress feature."""
    require_exact_training(train, "fit_et_walker")
    exact = train.loc[train["is_exact"]].reset_index(drop=True).copy()
    exact["walker_log_stress"] = walker_log_stress(exact, gamma)
    valid_aug = valid.copy()
    valid_aug["walker_log_stress"] = walker_log_stress(valid_aug, gamma)

    def new_pipe() -> Pipeline:
        return Pipeline([
            ("preprocess", make_preprocessor(("walker_log_stress",))),
            ("model", ExtraTreesRegressor(
                n_estimators=trees, min_samples_leaf=min_samples_leaf,
                max_features=max_features, random_state=seed, n_jobs=TREE_JOBS,
            )),
        ])

    pipe = new_pipe()
    pipe.fit(exact, exact["logN"], model__sample_weight=training_weights(exact))
    return np.asarray(pipe.predict(valid_aug), float)


def nested_select_et_walker(train: pd.DataFrame, seed: int, trees: int
                            ) -> tuple[dict[str, float], pd.DataFrame]:
    """Give ET-Walker exactly the WB-PIML inner folds and candidate budget."""
    require_exact_training(train, "nested_select_et_walker")
    splits, _ = grouped_folds(train, min(4, train[PRIMARY_GROUP_COLUMN].nunique()), seed)
    rows = []
    for candidate_index, (gamma, min_leaf, max_features, eta_budget) in enumerate(NESTED_CANDIDATES):
        pieces = []
        for inner_fold, (fit_idx, valid_idx) in enumerate(splits, 1):
            fit = train.iloc[fit_idx].reset_index(drop=True)
            valid = train.iloc[valid_idx].reset_index(drop=True)
            mu = fit_et_walker(
                fit, valid,
                seed + 10000 * candidate_index + 100 * inner_fold + 1,
                min(trees, 300), gamma, min_leaf, max_features,
            )
            part = valid[["row_id", "source_id", "campaign_id", "record_equivalence_id",
                          "logN", "is_exact", "is_runout"]].copy()
            part["mu"] = mu
            part["inner_fold"] = inner_fold
            pieces.append(part)
        oof = pd.concat(pieces, ignore_index=True)
        exact = oof.loc[oof["is_exact"]].copy()
        exact["pred_logN"] = exact["mu"]
        exact_score = sb_metrics([g.reset_index(drop=True) for _, g in exact.groupby(PRIMARY_GROUP_COLUMN)])["SB_RMSE"]
        censor_score = censored_hinge_rmse(oof)
        scale = calibrate_aft_scale(oof)
        rows.append({
            "candidate_index": candidate_index, "gamma": gamma,
            "min_samples_leaf": min_leaf, "max_features": max_features,
            "eta_budget_slot": eta_budget, "inner_SB_RMSE": exact_score,
            "inner_runout_hinge_RMSE": censor_score,
            "inner_joint_censored_NLL": scale["joint_censored_NLL"],
            "within_RMSE_tolerance": False,
        })
    candidates = pd.DataFrame(rows)
    best_rmse = float(candidates["inner_SB_RMSE"].min())
    candidates["within_RMSE_tolerance"] = (
        candidates["inner_SB_RMSE"] <= best_rmse + RMSE_TIE_TOLERANCE
    )
    best = choose_rmse_first(candidates)
    candidates = candidates.sort_values(
        ["within_RMSE_tolerance", "inner_SB_RMSE", "inner_joint_censored_NLL"],
        ascending=[False, True, True],
    ).reset_index(drop=True)
    chosen = {
        "candidate_index": int(best["candidate_index"]),
        "gamma": float(best["gamma"]),
        "min_samples_leaf": int(best["min_samples_leaf"]),
        "max_features": float(best["max_features"]),
        "eta_budget_slot": float(best["eta_budget_slot"]),
        "inner_SB_RMSE": float(best["inner_SB_RMSE"]),
        "inner_runout_hinge_RMSE": float(best["inner_runout_hinge_RMSE"]),
        "inner_joint_censored_NLL": float(best["inner_joint_censored_NLL"]),
        "RMSE_tolerance": RMSE_TIE_TOLERANCE,
    }
    return chosen, candidates


def external_predictions(train: pd.DataFrame, valid: pd.DataFrame, seed: int, trees: int,
                         et_min_samples_leaf: int = 3,
                         et_max_features: float = 0.85,
                         walker_gamma: float = GAMMA,
                         walker_min_samples_leaf: int = 3,
                         walker_max_features: float = 0.85) -> dict[str, np.ndarray]:
    require_exact_training(train, "external_predictions")
    exact = train.loc[train["is_exact"]].reset_index(drop=True)
    models = {
        "RF": RandomForestRegressor(n_estimators=trees, min_samples_leaf=3, max_features=0.75, random_state=seed, n_jobs=TREE_JOBS),
        "ExtraTrees": ExtraTreesRegressor(n_estimators=trees, min_samples_leaf=et_min_samples_leaf,
                                           max_features=et_max_features, random_state=seed + 1, n_jobs=TREE_JOBS),
        "GBDT": GradientBoostingRegressor(n_estimators=180, learning_rate=0.025, max_depth=2, min_samples_leaf=4, loss="huber", random_state=seed + 2),
        "SVR": SVR(C=2.0, epsilon=0.15, gamma="scale"),
    }
    output = {}
    for name, estimator in models.items():
        pipe = Pipeline([("preprocess", clone(make_preprocessor())), ("model", clone(estimator))])
        pipe.fit(exact, exact["logN"], model__sample_weight=training_weights(exact))
        output[name] = pipe.predict(valid)
    output["ET-Walker"] = fit_et_walker(
        train, valid, seed + 1, trees, walker_gamma,
        walker_min_samples_leaf, walker_max_features,
    )
    output["ML-Ens"] = np.mean(
        np.column_stack([output[name] for name in ["RF", "ExtraTrees", "GBDT", "SVR"]]),
        axis=1,
    )
    return output


def composite_score(error: np.ndarray, y: np.ndarray) -> float:
    q05, q95 = np.percentile(y, [5, 95]) if len(y) > 1 else (0.0, 1.0)
    scale = max(float(q95 - q05), LOG5)
    absolute = np.abs(error)
    return float(np.mean([
        np.clip(np.sqrt(np.mean(error**2)) / scale, 0, 1),
        np.clip(np.mean(absolute) / scale, 0, 1),
        np.mean(absolute > math.log10(3)), np.mean(absolute > LOG5),
        np.clip(abs(np.mean(error)) / scale, 0, 1),
    ]))


def _prepare_sb_groups(groups):
    prepared = []
    for group in groups:
        y = group["logN"].to_numpy(float)
        pred = group["pred_logN"].to_numpy(float)
        if "record_equivalence_id" in group.columns:
            equivalence_size = group.groupby("record_equivalence_id")[
                "record_equivalence_id"
            ].transform("size").to_numpy(float)
            weight = 1.0 / (
                max(group["record_equivalence_id"].nunique(), 1) * equivalence_size
            )
        else:
            weight = np.full(len(group), 1.0 / len(group))
        prepared.append((y, pred, weight, composite_score(pred - y, y)))
    return prepared


def _sb_metrics_prepared(groups):
    y = np.concatenate([g[0] for g in groups])
    pred = np.concatenate([g[1] for g in groups])
    weight = np.concatenate([g[2] for g in groups])
    error = pred - y
    mean_y = float(np.sum(weight * y) / np.sum(weight))
    denominator = float(np.sum(weight * (y - mean_y) ** 2))
    mse = float(np.sum(weight * error**2) / np.sum(weight))
    return {
        "SB_RMSE": math.sqrt(mse),
        "SB_MAE": float(np.sum(weight * np.abs(error)) / np.sum(weight)),
        "SB_R2": float(1.0 - np.sum(weight * error**2) / denominator) if denominator > 0 else np.nan,
        "SB_F5": float(np.sum(weight * (np.abs(error) <= LOG5)) / np.sum(weight)),
        "SB_C_star": float(np.mean([g[3] for g in groups])),
    }


def _sb_rmse_prepared(groups):
    y = np.concatenate([g[0] for g in groups])
    pred = np.concatenate([g[1] for g in groups])
    weight = np.concatenate([g[2] for g in groups])
    error = pred - y
    mse = float(np.sum(weight * error**2) / np.sum(weight))
    return math.sqrt(mse)



def sb_metrics(groups: list[pd.DataFrame]) -> dict[str, float]:
    """Return campaign-balanced point metrics, reported with the SB_* prefix."""
    return _sb_metrics_prepared(_prepare_sb_groups(groups))



def metric_summary(predictions: pd.DataFrame, bootstrap: int, seed: int) -> pd.DataFrame:
    rng = np.random.default_rng(seed)
    rows = []
    exact = predictions.loc[predictions["is_exact"]].copy()
    for model in legacy_first_models(exact["model"]):
        model_df = exact.loc[exact["model"].eq(model)]
        groups = _prepare_sb_groups([
            g.reset_index(drop=True) for _, g in model_df.groupby(PRIMARY_GROUP_COLUMN)
        ])
        point = _sb_metrics_prepared(groups)
        boots = {key: [] for key in point}
        for _ in range(bootstrap):
            sampled = [groups[i] for i in rng.integers(0, len(groups), len(groups))]
            values = _sb_metrics_prepared(sampled)
            for key, value in values.items():
                boots[key].append(value)
        row = {"model": model, "n_exact": len(model_df),
               "n_campaigns": len(groups), "n_sources": model_df["source_id"].nunique(), **point}
        for key, values in boots.items():
            row[f"{key}_lo"] = float(np.quantile(values, 0.025))
            row[f"{key}_hi"] = float(np.quantile(values, 0.975))
        rows.append(row)
    return pd.DataFrame(rows).sort_values("SB_RMSE").reset_index(drop=True)



def paired_effects(predictions: pd.DataFrame, bootstrap: int, seed: int,
                   comparator: str = "ML-Ens") -> pd.DataFrame:
    exact = predictions.loc[predictions["is_exact"] & predictions["model"].isin(["WB-PIML", comparator])].copy()
    wide = exact.pivot(
        index=["campaign_id", "source_id", "record_equivalence_id", "row_id", "logN"],
        columns="model", values="pred_logN"
    ).reset_index()
    groups = [g.reset_index(drop=True) for _, g in wide.groupby(PRIMARY_GROUP_COLUMN)]
    prepared = {
        model: _prepare_sb_groups([
            g[["record_equivalence_id", "logN", model]].rename(columns={model: "pred_logN"})
            for g in groups
        ])
        for model in ["WB-PIML", comparator]
    }

    def effect(indices):
        metrics = {
            model: _sb_metrics_prepared([prepared[model][i] for i in indices])
            for model in ["WB-PIML", comparator]
        }
        return {
            "RMSE_gain": metrics[comparator]["SB_RMSE"] - metrics["WB-PIML"]["SB_RMSE"],
            "MAE_gain": metrics[comparator]["SB_MAE"] - metrics["WB-PIML"]["SB_MAE"],
            "R2_gain": metrics["WB-PIML"]["SB_R2"] - metrics[comparator]["SB_R2"],
            "F5_gain": metrics["WB-PIML"]["SB_F5"] - metrics[comparator]["SB_F5"],
            "C_star_gain": metrics[comparator]["SB_C_star"] - metrics["WB-PIML"]["SB_C_star"],
        }

    point = effect(range(len(groups)))
    rng = np.random.default_rng(seed)
    boot = {key: [] for key in point}
    for _ in range(bootstrap):
        indices = rng.integers(0, len(groups), len(groups))
        values = effect(indices)
        for key, value in values.items():
            boot[key].append(value)
    return pd.DataFrame([{"comparator": comparator, "effect": key, "estimate": value,
                          "CI95_lo": float(np.quantile(boot[key], 0.025)),
                          "CI95_hi": float(np.quantile(boot[key], 0.975)),
                          "positive_is_better": True} for key, value in point.items()])



def uq_summary(predictions: pd.DataFrame, bootstrap: int, seed: int) -> pd.DataFrame:
    exact = predictions.loc[predictions["is_exact"] & predictions["model"].eq("WB-PIML")].copy()
    rng = np.random.default_rng(seed)
    rows = []
    for level in [80, 90]:
        low, high = f"lower{level}", f"upper{level}"
        exact["covered"] = (exact["logN"] >= exact[low]) & (exact["logN"] <= exact[high])
        exact["width"] = exact[high] - exact[low]
        alpha = 1.0 - level / 100.0
        exact["interval_score"] = exact["width"] + 2 / alpha * (exact[low] - exact["logN"]).clip(lower=0) + 2 / alpha * (exact["logN"] - exact[high]).clip(lower=0)
        groups = [g.reset_index(drop=True) for _, g in exact.groupby(PRIMARY_GROUP_COLUMN)]
        prepared = [
            (g["covered"].mean(), g["width"].mean(), g["interval_score"].mean())
            for g in groups
        ]

        def values(sampled):
            return {
                "coverage": float(np.mean([g[0] for g in sampled])),
                "mean_width": float(np.mean([g[1] for g in sampled])),
                "interval_score": float(np.mean([g[2] for g in sampled])),
            }

        point = values(prepared)
        boot = {key: [] for key in point}
        for _ in range(bootstrap):
            sampled = [prepared[i] for i in rng.integers(0, len(groups), len(groups))]
            for key, value in values(sampled).items():
                boot[key].append(value)
        row = {"nominal": level / 100.0, "n_exact": len(exact),
               "n_campaigns": exact["campaign_id"].nunique(),
               "n_sources": exact["source_id"].nunique(),
               "interval_completeness": float(np.mean(np.isfinite(exact[low]) & np.isfinite(exact[high]))),
               "calibration": "development grouped-OOF absolute residuals", **point}
        for key, values_ in boot.items():
            row[f"{key}_lo"] = float(np.quantile(values_, 0.025))
            row[f"{key}_hi"] = float(np.quantile(values_, 0.975))
        rows.append(row)
    return pd.DataFrame(rows)



def runout_summary(predictions: pd.DataFrame) -> pd.DataFrame:
    rows = []
    audited_models = ([s.key for s in PHYSICAL_SPECS]
                      + [PROPOSED_MODEL, "ET-Walker", "ExtraTrees"]
                      + list(CENSORED_MODELS) + ["ML-Strong"] + list(SIMPLE_CONTROL_MODELS))
    runout = predictions.loc[
        predictions["is_runout"] & predictions["model"].isin(audited_models)
    ].copy()
    for model, group in runout.groupby("model"):
        weight = analysis_weights(group)
        finite_sigma = np.isfinite(group["sigma"].to_numpy(float))
        nll_value = np.nan
        if finite_sigma.all():
            if model == "Weibull-AFT":
                z = ((group["logN"].to_numpy(float) - group["pred_logN"].to_numpy(float))
                     / group["sigma"].to_numpy(float))
                loss = np.exp(np.clip(z, -30.0, 30.0))
            else:
                z = ((group["pred_logN"].to_numpy(float) - group["logN"].to_numpy(float))
                     / group["sigma"].to_numpy(float))
                loss = -scipy_special.log_ndtr(z)
            nll_value = float(np.sum(weight * loss) / np.sum(weight))
        below = (group["pred_logN"] < group["logN"]).astype(float)
        group_rates = below.groupby(group[PRIMARY_GROUP_COLUMN]).mean()
        rows.append({"model": model, "n_runout": len(group),
                     "n_campaigns": group["campaign_id"].nunique(),
                     "n_sources": group["source_id"].nunique(),
                     "SB_runout_NLL": nll_value,
                     "SB_point_below_runout_limit_rate": float(group_rates.mean()),
                     "runouts_in_location_training": bool(group["used_for_training"].all()),
                     "runouts_in_candidate_selection": bool(
                         group["used_for_candidate_selection"].all()
                     ),
                     "used_as_censor_feasibility_reference": bool(
                         group["used_as_censor_feasibility_reference"].all()
                     ),
                     "runouts_in_model_selection_or_reference": bool(
                         group["used_for_selection"].all()
                     ),
                     "interpretation": "outer-test right-censored audit; stopping cycle is a lower bound, not a point failure life"})
    return pd.DataFrame(rows).sort_values("SB_runout_NLL", na_position="last")


def source_gain_table(predictions: pd.DataFrame) -> pd.DataFrame:
    exact = predictions.loc[
        predictions["is_exact"] & predictions["model"].isin(["ML-Ens", "WB-PIML"])
    ]
    wide = exact.pivot(
        index=["campaign_id", "source_id", "record_equivalence_id", "row_id", "logN"],
        columns="model", values="pred_logN"
    ).reset_index()
    rows = []
    for source, group in wide.groupby("source_id"):
        equivalence_size = group.groupby("record_equivalence_id")[
            "record_equivalence_id"
        ].transform("size").to_numpy(float)
        weight = 1.0 / (
            max(group["record_equivalence_id"].nunique(), 1) * equivalence_size
        )
        ml_error = (group["ML-Ens"] - group["logN"]).to_numpy(float)
        piml_error = (group["WB-PIML"] - group["logN"]).to_numpy(float)
        rmse_ml = float(np.sqrt(np.sum(weight * ml_error**2) / np.sum(weight)))
        rmse_piml = float(np.sqrt(np.sum(weight * piml_error**2) / np.sum(weight)))
        rows.append({"source_id": source, "n_exact": len(group), "RMSE_ML_Ens": rmse_ml,
                     "RMSE_WB_PIML": rmse_piml, "RMSE_gain": rmse_ml - rmse_piml})
    table = pd.DataFrame(rows).sort_values("RMSE_gain", ascending=False).reset_index(drop=True)
    label_map = {
        source: f"S{index:02d}"
        for index, source in enumerate(sorted(table["source_id"].astype(str).unique()), 1)
    }
    table.insert(1, "plot_label", table["source_id"].astype(str).map(label_map))
    return table


def ablation_effects(predictions: pd.DataFrame, bootstrap: int, seed: int) -> pd.DataFrame:
    models = ["WB", "WB-CD", "WB-M", "WB-Ox", "ExtraTrees", "ET-Walker",
              "WB-PIML", "WB-PIML-M", "WB-PIML-Ox", "ML-Ens"]
    exact = predictions.loc[predictions["is_exact"] & predictions["model"].isin(models)]
    wide = exact.pivot(
        index=["campaign_id", "source_id", "record_equivalence_id", "row_id", "logN"],
        columns="model", values="pred_logN"
    ).reset_index()
    groups = [g.reset_index(drop=True) for _, g in wide.groupby(PRIMARY_GROUP_COLUMN)]
    prepared = {
        model: _prepare_sb_groups([
            g[["record_equivalence_id", "logN", model]].rename(columns={model: "pred_logN"})
            for g in groups
        ])
        for model in models
    }

    def rmse(indices, model):
        return _sb_rmse_prepared([prepared[model][i] for i in indices])

    rng = np.random.default_rng(seed)
    rows = []
    for model in models:
        if model == PROPOSED_MODEL:
            continue
        point = rmse(range(len(groups)), model) - rmse(range(len(groups)), PROPOSED_MODEL)
        draws = []
        for _ in range(bootstrap):
            indices = rng.integers(0, len(groups), len(groups))
            draws.append(rmse(indices, model) - rmse(indices, PROPOSED_MODEL))
        rows.append({"comparison": f"{PROPOSED_MODEL} gain over {model}", "RMSE_gain": point,
                     "CI95_lo": float(np.quantile(draws, 0.025)),
                     "CI95_hi": float(np.quantile(draws, 0.975)),
                     "positive_favors_WB_PIML": True})
    return pd.DataFrame(rows)
