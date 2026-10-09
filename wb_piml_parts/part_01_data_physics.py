# ===== REVISED PRIMARY CODE BEGIN =====


from __future__ import annotations
from scipy.linalg import cho_factor, cho_solve
from sklearn.gaussian_process.kernels import Matern
import itertools

import argparse
import pickle
from collections import OrderedDict
from functools import wraps
import hashlib
import json
import math
import os
import platform
import tempfile
import time
import warnings
import zipfile
from dataclasses import dataclass, replace
from io import BytesIO
from pathlib import Path

os.environ.setdefault("MPLCONFIGDIR", os.path.join(os.environ.get("TEMP", "/tmp"), "piml_mpl"))

import matplotlib.pyplot as plt
from matplotlib.transforms import Bbox
import numpy as np
import pandas as pd
import scipy
import sklearn
from scipy.optimize import minimize
from joblib import Parallel, delayed, parallel_config
from threadpoolctl import threadpool_limits
import scipy.special as scipy_special
from sklearn.base import clone
from sklearn.compose import ColumnTransformer
from sklearn.ensemble import ExtraTreesRegressor, GradientBoostingRegressor, RandomForestRegressor
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import OneHotEncoder, StandardScaler
from sklearn.svm import SVR

try:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    Alignment = Font = PatternFill = get_column_letter = None

try:
    from PIL import Image
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt
except ImportError:
    Image = Presentation = RGBColor = PP_ALIGN = Inches = Pt = None


WORKERS = 8
TREE_JOBS = 1
BLAS_THREADS = 1

SEED = 20260824
GAMMA = 0.85
DAMAGE_CAP = 0.40
ARRHENIUS_REFERENCE_C = 800.0
ARRHENIUS_SCALE = 1000.0
LOG5 = math.log10(5.0)
SHEET = "Plot_Data_Verified"
EXPECTED_FILE = "data.xlsx"
EXPECTED_VERIFIED_ROWS = 224
EXPECTED_EXACT_ROWS = 158
EXPECTED_RUNOUT_ROWS = 66
DEFAULT_OUTPUT_DIR = Path(__file__).resolve().parent / "WB_PIML_results"
PPTX_OUTPUT_NAME = "WB_PIML_figures.pptx"
PRIMARY_GROUP_COLUMN = "campaign_id"
PROXIES = ["Dm", "Di", "Df", "Dox"]
TRAD_NUMERIC = ["logS", "R", "Tn", "logf", "logUTS", "Sa", "Smean"]
TRAD_CATEGORICAL = ["architecture", "architecture_detail", "env_class"]


RESIDUAL_INPUT_POLICY = "environment_shared_residual_v1_20261008"
RESIDUAL_NUMERIC = tuple(TRAD_NUMERIC)
RESIDUAL_CATEGORICAL = ("architecture", "architecture_detail")
RESIDUAL_PHYSICS_NUMERIC = ("mechanical_prediction", "walker_stress")
NESTED_CANDIDATES = (
    
    (0.75, 2, 0.65, 0.50),
    (0.75, 3, 0.85, 1.00),
    (0.75, 5, 1.00, 0.50),
    (0.85, 2, 0.85, 1.00),
    (0.85, 3, 0.65, 0.50),
    (0.85, 3, 0.85, 0.00),
    (0.85, 3, 0.85, 0.50),
    (0.85, 3, 0.85, 1.00),
    (0.85, 5, 0.85, 1.00),
    (0.95, 2, 0.85, 1.00),
    (0.95, 3, 0.85, 0.50),
    (0.95, 3, 1.00, 1.00),
)
STRUCTURAL_GAMMAS = (.75, .85, .95)
ENVIRONMENT_STRENGTH_GRID = (0., .1, 1.)
RESIDUAL_GAIN_GRID = (.75, 1.)
HYBRID_FAMILIES = ("ET_residual", "physics_feature_ET_residual", "Matern_residual")
TEMPERATURE_POLICIES = ("soft250",)
def structural_candidate_grid():
    hybrids = [dict(pool="hybrid", family=family, gamma=gamma, strength=strength, eta=eta)
               for family, gamma, strength, eta in itertools.product(HYBRID_FAMILIES, STRUCTURAL_GAMMAS, ENVIRONMENT_STRENGTH_GRID, RESIDUAL_GAIN_GRID)]
    controls = [dict(pool="control", family="ET", leaf=leaf, max_features=mf)
                for leaf, mf in itertools.product((2, 3, 5, 8, 12, 16), (.65, .85, 1.))]
    controls += [dict(pool="control", family="GBDT", depth=depth, leaf=leaf, learning_rate=lr)
                 for depth, leaf, lr in itertools.product((1, 2, 3), (3, 8, 12), (.03, .07))]
    controls += [dict(pool="control", family="Matern", length_scale=length, alpha=alpha)
                 for length, alpha in itertools.product((1., 2., 4.), (.1, .3, 1., 3., 10., 30.))]
    controls += [dict(pool="control", family="fixed_ML_Ens")]
    controls += [dict(pool="control", family="physics_feature_ET", gamma=gamma, strength=strength)
                 for gamma, strength in itertools.product(STRUCTURAL_GAMMAS, ENVIRONMENT_STRENGTH_GRID)]
    return [dict(candidate_id=i, **config) for i, config in enumerate(hybrids + controls)]


STRUCTURAL_CANDIDATES = tuple(structural_candidate_grid())
HYBRID_CANDIDATES = tuple(c for c in STRUCTURAL_CANDIDATES if c["pool"] == "hybrid")
CONTROL_CANDIDATES = tuple(c for c in STRUCTURAL_CANDIDATES if c["pool"] == "control")
HYBRID_BASES = tuple({k:v for k,v in c.items() if k not in ("eta", "candidate_id")} for c in HYBRID_CANDIDATES[::2])

PHYSICS_MIX_GRID = (1.00,)
CALIBRATION_MODES = ("none",)
RMSE_TIE_TOLERANCE = 0.001
NEAR_TIE_RMSE = 0.020

RESIDUAL_TEMPERATURE_MATCH_ATOL_C = 1e-9
CENSORED_MODELS = ("Lognormal-AFT", "Weibull-AFT")
REPEATED_VALIDATION_SEEDS = tuple(SEED + 101 * index for index in range(10))


AUDIT_CONTEXT: dict[str, object] = {}
ACTIVE_DATA_PATH: Path | None = None


@dataclass(frozen=True)
class PhysicalSpec:
    key: str
    channels: tuple[str, ...]
    mode: str = "linear"


@dataclass(frozen=True)
class CompetingDamageConfig:


    intercept_half_width: float = 0.75
    walker_half_width: float = 3.0
    walker_min: float = 0.05
    walker_max: float = 20.0
    environment_ratio_min: float = -8.0
    environment_ratio_max: float = 4.0
    environment_stress_min: float = 0.0
    environment_stress_max: float = 10.0
    arrhenius_min: float = 0.0
    arrhenius_max: float = 10.0
    ridge_strength: float = 0.01
    environment_strength: float = 0.0


DEFAULT_COMPETING_CONFIG = CompetingDamageConfig()


CLASSIC_WB_SPEC = PhysicalSpec("WB", ())
COMPETING_WB_SPEC = PhysicalSpec("WB-CD", (), "competing_damage")
PHYSICAL_SPECS = (
    CLASSIC_WB_SPEC,
    COMPETING_WB_SPEC,
    PhysicalSpec("WB-M", ("M",)),
    PhysicalSpec("WB-Ox", ("Ox",)),
)
EXTERNAL_MODELS = ("RF", "ExtraTrees", "ET-Walker", "GBDT", "SVR", "ML-Ens", "ML-Strong")
PRIMARY_COMPARATORS = ("ML-Ens", "ExtraTrees", "ET-Walker", "ML-Strong")
HYBRID_SPECS = (
    ("WB-PIML-Anchor", COMPETING_WB_SPEC),
    ("WB-PIML-M", PhysicalSpec("WB-M", ("M",))),
    ("WB-PIML-Ox", PhysicalSpec("WB-Ox", ("Ox",))),
)
PROPOSED_MODEL = "WB-PIML"
SIMPLE_CONTROL_MODELS = ("WB-Residual", "Basquin-Residual")


def legacy_first_models(models):

    names = set(models)
    return sorted(names - set(SIMPLE_CONTROL_MODELS)) + [
        name for name in SIMPLE_CONTROL_MODELS if name in names
    ]


def parse_bool(series: pd.Series, name: str = "boolean field") -> pd.Series:
    if pd.api.types.is_bool_dtype(series):
        return series.astype(bool)
    mapped = series.astype(str).str.strip().str.lower().map(
        {"true": True, "1": True, "yes": True, "y": True,
         "false": False, "0": False, "no": False, "n": False}
    )
    if mapped.isna().any():
        raise ValueError(f"{name} contains unrecognized values")
    return mapped.astype(bool)


def resolve_data(argument: str | None) -> Path:
    here = Path(__file__).resolve().parent
    candidates = [] if argument is None else [Path(argument)]
    candidates += [
        here / EXPECTED_FILE,
        Path.cwd() / EXPECTED_FILE,
        here / "upload" / EXPECTED_FILE,
    ]
    for path in candidates:
        if path.exists():
            return path.resolve()
    raise FileNotFoundError(f"Cannot find {EXPECTED_FILE}; pass --data")


def require_exact_training(frame: pd.DataFrame, context: str) -> pd.DataFrame:

    if frame.empty:
        raise ValueError(f"{context}: empty training frame")
    if "is_runout" not in frame or "is_exact" not in frame:
        raise KeyError(f"{context}: missing exact/runout labels")
    if not frame["is_exact"].eq(~frame["is_runout"]).all():
        raise AssertionError(f"{context}: exact/runout labels are not complementary")
    if not frame["is_exact"].any():
        raise AssertionError(f"{context}: at least one exact failure is required")
    if "sample_weight" in frame.columns:
        supplied = pd.to_numeric(frame["sample_weight"], errors="coerce").to_numpy(float)
        if not np.isfinite(supplied).all() or not np.allclose(supplied, 1.0):
            raise ValueError(f"{context}: every analysis record must have weight 1")
    return frame


def sigmoid(x: np.ndarray) -> np.ndarray:
    return 1.0 / (1.0 + np.exp(-np.clip(x, -40.0, 40.0)))


def apply_physics_drivers(df: pd.DataFrame, water_factor: float | None = None,
                          arrhenius_reference_c: float = ARRHENIUS_REFERENCE_C,
                          arrhenius_scale: float = ARRHENIUS_SCALE) -> pd.DataFrame:


    if (water_factor is not None and (not np.isfinite(water_factor) or water_factor <= 0)) or arrhenius_scale <= 0:
        raise ValueError("water_factor and arrhenius_scale must be positive")
    result = df.copy()
    env = result["env_class"].astype(str).str.lower()
    inert = env.str.contains("vacuum|inert|argon", regex=True, na=False)
    water = env.str.contains("water|steam", regex=True, na=False)
    result["environment_exposure"] = np.where(
        inert.to_numpy(bool), 0.0,
        np.where(water.to_numpy(bool), 1.0 if water_factor is None else float(water_factor), 1.0),
    )
    result["fit_water_multiplier"] = water_factor is None
    temperature_k = result["T_C"].to_numpy(float) + 273.15
    reference_k = float(arrhenius_reference_c) + 273.15
    if reference_k <= 0 or np.any(temperature_k <= 0):
        raise ValueError("Arrhenius temperatures must be above absolute zero")
    result["arrhenius_temperature_drive"] = float(arrhenius_scale) * (
        1.0 / reference_k - 1.0 / temperature_k
    )
    return result


def load_data(path: Path) -> pd.DataFrame:
    if SHEET not in pd.ExcelFile(path).sheet_names:
        raise KeyError(f"Missing sheet {SHEET!r}")
    configure_audit_context(path)
    df = pd.read_excel(path, sheet_name=SHEET).copy()
    required = {
        "row_id", "source_id", "architecture", "arch_group", "env_class", "R", "T_C",
        "frequency_Hz", "Nf_cycles", "UTS_MPa", "stress_level_sigma_max_over_UTS",
        "sigma_max_MPa_before_normalization", "is_runout", "is_exact_failure",
        "record_equivalence_id", "campaign_id", "training_role", "sample_weight",
        "event_observed", "censoring_type", "likelihood_weight", "point_metric_eligible",
    }
    missing = sorted(required - set(df.columns))
    if missing:
        raise KeyError(f"Missing required columns: {missing}")
    df["row_id"] = pd.to_numeric(df["row_id"], errors="raise").astype(int)
    if df["row_id"].duplicated().any():
        raise ValueError("row_id must be unique")
    df["source_id"] = df["source_id"].fillna("").astype(str).str.strip()
    if df["source_id"].eq("").any():
        raise ValueError("source_id contains blanks")
    df["architecture_detail"] = df["architecture"].fillna("unknown").astype(str).str.strip().str.lower()
    df["architecture"] = (df["arch_group"].fillna(df["architecture_detail"]).fillna("unknown")
                          .astype(str).str.strip().str.lower())
    df["env_class"] = df["env_class"].fillna("unknown").astype(str).str.strip().str.lower()
    if "temp_bin" in df.columns:
        df["temp_bin"] = df["temp_bin"].fillna("unknown").astype(str).str.strip().str.lower()
    else:
        
        df["temp_bin"] = pd.to_numeric(df["T_C"], errors="coerce").map(lambda value: f"{value:g}c")
    for col in ["R", "T_C", "frequency_Hz", "Nf_cycles", "UTS_MPa",
                "stress_level_sigma_max_over_UTS", "sigma_max_MPa_before_normalization"]:
        df[col] = pd.to_numeric(df[col], errors="coerce")
    numeric = ["R", "T_C", "frequency_Hz", "Nf_cycles", "UTS_MPa",
               "stress_level_sigma_max_over_UTS", "sigma_max_MPa_before_normalization"]
    if df[numeric].isna().any().any() or not np.isfinite(df[numeric].to_numpy(float)).all():
        raise ValueError("Required numeric columns contain missing/non-numeric values")
    
    boundary = df["T_C"].eq(800) & df["temp_bin"].isin(["500-800", "800-1000"])
    df.loc[boundary, "temp_bin"] = "800-1000"
    temperature_groups = df.groupby("T_C")["temp_bin"].nunique()
    if temperature_groups.gt(1).any():
        raise ValueError("Each temperature must belong to one temp_bin; conflicting temperatures: "
                         f"{temperature_groups.index[temperature_groups.gt(1)].tolist()}")
    if (df[["frequency_Hz", "Nf_cycles", "UTS_MPa", "stress_level_sigma_max_over_UTS"]] <= 0).any().any():
        raise ValueError("Frequency, life, UTS and normalized stress must be positive")
    if (df["R"] >= 1.0).any():
        raise ValueError("Walker correction requires every stress ratio R < 1")
    df["is_runout"] = parse_bool(df["is_runout"], "is_runout")
    supplied_exact = parse_bool(df["is_exact_failure"], "is_exact_failure")
    mismatch = supplied_exact.ne(~df["is_runout"])
    if mismatch.any():
        bad_rows = df.loc[mismatch, "row_id"].astype(int).tolist()
        raise ValueError(
            "is_exact_failure must equal ~is_runout; mismatched row_id values: "
            f"{bad_rows[:20]}"
        )
    df["supplied_is_runout"] = df["is_runout"].astype(bool)
    df["supplied_is_exact_failure"] = supplied_exact.astype(bool)
    df["label_override_applied"] = False
    df["label_review_status"] = "accepted_from_Plot_Data_Verified"
    if "evidence_source" in df.columns:
        df["label_review_evidence"] = df["evidence_source"].fillna("").astype(str)
    else:
        df["label_review_evidence"] = "verified workbook label retained without code-side override"
    df["reviewed_is_exact_failure"] = ~df["is_runout"]
    df["is_exact"] = ~df["is_runout"]
    df["label_consistent"] = df["reviewed_is_exact_failure"].eq(~df["is_runout"])
    df["training_role"] = df["training_role"].fillna("").astype(str).str.strip()
    df["sample_weight"] = pd.to_numeric(df["sample_weight"], errors="coerce")
    expected_role = np.where(
        df["is_exact"], "exact_likelihood_training", "right_censored_likelihood_training"
    )
    role_mismatch = df["training_role"].ne(expected_role)
    if role_mismatch.any():
        bad = df.loc[role_mismatch, ["row_id", "training_role", "is_runout"]]
        raise ValueError(
            "training_role is inconsistent with the joint-likelihood contract: "
            f"{bad.to_dict('records')[:20]}"
        )
    df["likelihood_weight"] = pd.to_numeric(df["likelihood_weight"], errors="coerce")
    if not np.allclose(df["sample_weight"], 1.0) or not np.allclose(df["likelihood_weight"], 1.0):
        bad = df.loc[
            ~np.isclose(df["sample_weight"], 1.0) | ~np.isclose(df["likelihood_weight"], 1.0),
            ["row_id", "training_role", "sample_weight", "likelihood_weight"],
        ]
        raise ValueError(f"Every unique record must have unit likelihood weight: {bad.to_dict('records')[:20]}")
    event = pd.to_numeric(df["event_observed"], errors="coerce")
    point = pd.to_numeric(df["point_metric_eligible"], errors="coerce")
    expected_event = df["is_exact"].astype(int)
    expected_censoring = np.where(df["is_exact"], "exact", "right_censored")
    if not event.eq(expected_event).all() or not point.eq(expected_event).all():
        raise ValueError("event_observed/point_metric_eligible are inconsistent with exact/runout labels")
    if not df["censoring_type"].astype(str).eq(expected_censoring).all():
        raise ValueError("censoring_type is inconsistent with exact/runout labels")
    if "include_in_verified_model" in df.columns:
        included = parse_bool(df["include_in_verified_model"], "include_in_verified_model")
        if not included.all():
            bad = df.loc[~included, "row_id"].astype(int).tolist()
            raise ValueError(
                "Plot_Data_Verified may contain only retained unique records; "
                f"excluded row_id values were found: {bad[:20]}"
            )
    df["logN"] = np.log10(df["Nf_cycles"].clip(lower=1.0))
    df["S"] = df["stress_level_sigma_max_over_UTS"]
    df["logS"] = np.log10(df["S"].clip(lower=1e-8))
    df["logf"] = np.log10(df["frequency_Hz"].clip(lower=1e-8))
    df["logUTS"] = np.log10(df["UTS_MPa"].clip(lower=1e-8))
    df["Tn"] = (df["T_C"] - 25.0) / 1000.0
    df["Sa"] = 0.5 * df["S"] * np.clip(1.0 - df["R"], 0.0, 2.0)
    df["Smean"] = 0.5 * df["S"] * (1.0 + df["R"])

    df = apply_physics_drivers(df)
    env = df["env_class"].str.lower()
    inert = env.str.contains("vacuum|inert|argon", regex=True, na=False)
    water = env.str.contains("water|steam", regex=True, na=False)
    thermal = sigmoid((df["T_C"].to_numpy(float) - 800.0) / 200.0)
    slow = sigmoid(-df["logf"].to_numpy(float))
    oxidation = (~inert).to_numpy(float) * thermal * (0.5 + 0.5 * slow)
    oxidation = np.maximum(oxidation, water.to_numpy(float) * thermal)
    df["oxidation"] = np.clip(oxidation, 0.0, 1.0)
    stress = np.clip(df["S"].to_numpy(float), 0.0, 1.0)
    df["Dm"] = stress
    df["Di"] = np.clip(stress * (0.60 + 0.40 * df["oxidation"]), 0.0, 1.0)
    df["Df"] = np.clip(stress**2, 0.0, 1.0)
    df["Dox"] = df["oxidation"]
    df["M"] = df[["Dm", "Di", "Df"]].mean(axis=1)
    df["Ox"] = df["Dox"]
    df["Q"] = df[PROXIES].mean(axis=1)
    series_columns = ["source_id", "architecture_detail", "env_class", "T_C", "R", "UTS_MPa"]
    df["series_group_id"] = df[series_columns].astype(str).agg("|".join, axis=1)
    if "raw_Nf" in df.columns:
        raw_nf = df["raw_Nf"].fillna("").astype(str).str.strip().str.lower()
    else:
        raw_nf = pd.Series("", index=df.index)
    round_limit = np.isclose(df["Nf_cycles"], 1e6) | np.isclose(df["Nf_cycles"], 1e7)
    df["qa_exact_round_limit"] = df["is_exact"] & round_limit & ~raw_nf.str.contains("~|未断|runout|>", regex=True)
    df["qa_sstar_gt_1"] = df["S"] > 1.0
    return add_campaign_metadata(df.reset_index(drop=True))


def source_weights(df: pd.DataFrame) -> np.ndarray:
    count = df.groupby("source_id")["source_id"].transform("size").to_numpy(float)
    weight = 1.0 / count
    return weight / np.mean(weight)


def training_weights(df: pd.DataFrame) -> np.ndarray:


    require_exact_training(df, "training_weights")
    if "sample_weight" in df.columns:
        supplied = pd.to_numeric(df["sample_weight"], errors="coerce").to_numpy(float)
        if not np.isfinite(supplied).all() or not np.allclose(supplied, 1.0):
            bad = df.loc[
                ~np.isclose(supplied, 1.0), ["row_id", "sample_weight"]
            ].to_dict("records")
            raise ValueError(f"Every unique likelihood record must have weight 1: {bad[:20]}")
    return np.ones(len(df), dtype=float)


def joint_aft_nll(frame: pd.DataFrame, mu: np.ndarray, sigma: float,
                  weights: np.ndarray | None = None) -> dict[str, float]:

    require_exact_training(frame, "joint_aft_nll")
    if not np.isfinite(sigma) or sigma <= 0:
        raise ValueError("AFT sigma must be positive and finite")
    location = np.asarray(mu, float)
    if len(location) != len(frame) or not np.isfinite(location).all():
        raise ValueError("AFT location prediction is missing or non-finite")
    weight = training_weights(frame) if weights is None else np.asarray(weights, float)
    y = frame["logN"].to_numpy(float)
    exact_mask = frame["is_exact"].to_numpy(bool)
    runout_mask = ~exact_mask
    z = (y - location) / sigma
    losses = np.empty(len(frame), float)
    losses[exact_mask] = (
        0.5 * math.log(2.0 * math.pi) + math.log(sigma)
        + 0.5 * z[exact_mask] ** 2
    )
    
    losses[runout_mask] = -scipy_special.log_ndtr(-z[runout_mask])
    total = float(np.sum(weight * losses) / np.sum(weight))
    exact_nll = float(np.average(losses[exact_mask], weights=weight[exact_mask]))
    runout_nll = (
        float(np.average(losses[runout_mask], weights=weight[runout_mask]))
        if runout_mask.any() else np.nan
    )
    return {
        "joint_censored_NLL": total,
        "exact_density_NLL": exact_nll,
        "runout_survival_NLL": runout_nll,
    }


def analysis_weights(df: pd.DataFrame) -> np.ndarray:


    required = {"campaign_id", "record_equivalence_id"}
    if not required.issubset(df.columns):
        return source_weights(df)
    frame = df[["campaign_id", "record_equivalence_id"]].astype(str)
    pair_size = frame.groupby(
        ["campaign_id", "record_equivalence_id"]
    )["record_equivalence_id"].transform("size").to_numpy(float)
    equivalence_per_campaign = (
        frame.drop_duplicates()
        .groupby("campaign_id")["record_equivalence_id"].nunique()
    )
    n_equivalence = frame["campaign_id"].map(equivalence_per_campaign).to_numpy(float)
    weight = 1.0 / (pair_size * n_equivalence)
    return weight / np.mean(weight)


def grouped_folds(df: pd.DataFrame, n_splits: int, seed: int):

    require_exact_training(df, "grouped_folds")
    group_column = PRIMARY_GROUP_COLUMN
    summary = df.groupby(group_column).agg(n=("row_id", "size")).reset_index()
    n_splits = min(n_splits, len(summary))
    rng = np.random.default_rng(seed)
    summary["tie"] = rng.random(len(summary))
    summary = summary.sort_values(["n", "tie"], ascending=[False, True])
    target = float(summary["n"].sum()) / n_splits
    load = np.zeros(n_splits)
    assignment: dict[str, int] = {}
    for _, row in summary.iterrows():
        score = load / max(target, 1.0)
        fold = int(np.argmin(score))
        load[fold] += float(row["n"])
        assignment[str(row[group_column])] = fold
    fold_id = df[group_column].astype(str).map(assignment).to_numpy(int)
    splits = []
    for fold in range(n_splits):
        valid = np.flatnonzero(fold_id == fold)
        train = np.flatnonzero(fold_id != fold)
        if not len(train) or not len(valid):
            raise RuntimeError("Empty grouped fold")
        if set(df.iloc[train][group_column]) & set(df.iloc[valid][group_column]):
            raise AssertionError("Campaign leakage")
        if not df.iloc[train]["is_exact"].any() or not df.iloc[valid]["is_exact"].any():
            raise RuntimeError("Every outer fold must contain exact failures")
        splits.append((train, valid))
    return splits, fold_id + 1


def walker_log_stress(df: pd.DataFrame, gamma: float = GAMMA) -> np.ndarray:
    factor = np.clip(0.5 * (1.0 - df["R"].to_numpy(float)), 1e-6, None)
    return df["logS"].to_numpy(float) + gamma * np.log10(factor)


def physical_design(df: pd.DataFrame, spec: PhysicalSpec, gamma: float = GAMMA) -> np.ndarray:
    if spec.mode != "linear":
        raise ValueError(f"{spec.key} uses a nonlinear competing-damage equation")
    columns = [np.ones(len(df)), -walker_log_stress(df, gamma)]
    columns += [-df[channel].to_numpy(float) for channel in spec.channels]
    return np.column_stack(columns)


def _fit_linear_physical(frame: pd.DataFrame, spec: PhysicalSpec,
                         gamma: float) -> np.ndarray:

    require_exact_training(frame, f"_fit_linear_physical[{spec.key}]")
    x = physical_design(frame, spec, gamma)
    exact_mask = frame["is_exact"].to_numpy(bool)
    initial_coef = np.linalg.lstsq(
        x[exact_mask], frame.loc[exact_mask, "logN"].to_numpy(float), rcond=None
    )[0]
    initial_coef[1] = np.clip(initial_coef[1], 0.05, 20.0)
    if len(initial_coef) > 2:
        initial_coef[2:] = np.clip(initial_coef[2:], 0.0, DAMAGE_CAP)
    initial_sigma = np.clip(
        np.std(frame.loc[exact_mask, "logN"].to_numpy(float) - x[exact_mask] @ initial_coef),
        0.10, 3.0,
    )

    def loss(parameter: np.ndarray) -> float:
        coef = parameter[:-1]
        sigma = float(np.exp(parameter[-1]))
        nll = joint_aft_nll(frame, x @ coef, sigma)["joint_censored_NLL"]
        ridge = 1e-4 * float(np.sum(coef[2:] ** 2)) if len(coef) > 2 else 0.0
        return nll + ridge

    initial = np.r_[initial_coef, math.log(initial_sigma)]
    bounds = ([(None, None), (0.05, 20.0)]
              + [(0.0, DAMAGE_CAP)] * len(spec.channels)
              + [(math.log(0.08), math.log(4.0))])
    result = minimize(
        loss, initial, method="L-BFGS-B", bounds=bounds,
        options={"maxiter": 1200, "ftol": 1e-12, "gtol": 1e-8},
    )
    if not result.success or not np.isfinite(result.fun):
        
        retry_start = (
            np.asarray(result.x, float)
            if np.size(result.x) == np.size(initial)
            and np.all(np.isfinite(result.x))
            else initial
        )
        fallback = minimize(
            loss, retry_start, method="Powell", bounds=bounds,
            options={"maxiter": 6000, "xtol": 1e-9, "ftol": 1e-10},
        )
        if fallback.success and np.isfinite(fallback.fun):
            result = fallback
        else:
            raise RuntimeError(
                f"{spec.key} censored-likelihood optimization failed: "
                f"L-BFGS-B={result.message}; Powell={fallback.message}"
            )
    return np.asarray(result.x[:-1], float)


def _predict_competing_damage(df: pd.DataFrame, coef: np.ndarray,
                              gamma: float) -> np.ndarray:


    if len(coef) not in (5,6):
        raise ValueError("Competing-damage parameters: five fixed-weight or six learned-weight coefficients")
    intercept, b, log_k, m, q_t = np.asarray(coef, float)[:5]
    log_sw = walker_log_stress(df, gamma)
    log_d_mech = -intercept + b * log_sw

    exposure = df["environment_exposure"].to_numpy(float)
    active = exposure > 0.0
    log_d_env = np.full(len(df), -np.inf, dtype=float)
    log_d_env[active] = (
        log_d_mech[active]
        + log_k
        + np.log10(exposure[active]) + water_log_adjustment(df,coef)[active]
        + (q_t * df.loc[active, "arrhenius_temperature_drive"].to_numpy(float)
           / math.log(10.0))
        + m * log_sw[active]
        - df.loc[active, "logf"].to_numpy(float)
    )
    
    maximum = np.maximum(log_d_mech, log_d_env)
    log_total_damage = maximum + np.log10(
        np.power(10.0, log_d_mech - maximum)
        + np.power(10.0, log_d_env - maximum)
    )
    return -log_total_damage


def cache_physical_fit(function, validate, damage_cap, maxsize=2048):
    cache = OrderedDict()
    stats = {'hits': 0, 'misses': 0}
    gamma_default, config_default = function.__defaults__

    def array_token(values):
        array = np.asarray(values)
        content = (
            tuple(pickle.dumps(value, protocol=5) for value in array.flat)
            if array.dtype.hasobject else array.tobytes(order='C')
        )
        return array.dtype.str, array.shape, content

    def key_for(train, spec, gamma, competing_config):
        dtypes = tuple(map(str, train.dtypes))
        payload = (
            tuple(train.columns.names), array_token(train.columns),
            tuple(train.index.names), array_token(train.index),
            tuple((dtypes[i], array_token(train.iloc[:, i].to_numpy()))
                  for i in range(train.shape[1])),
            (spec.key, tuple(spec.channels), spec.mode), gamma,
            tuple(sorted(vars(competing_config).items())), damage_cap(),
        )
        return hashlib.sha256(pickle.dumps(payload, protocol=5)).digest()

    @wraps(function)
    def cached(train, spec, gamma=gamma_default, competing_config=config_default):
        validate(train, f"fit_physical[{spec.key}]")
        key = key_for(train, spec, gamma, competing_config)
        if key in cache:
            stats['hits'] += 1
            cache.move_to_end(key)
            return cache[key].copy()
        result = function(train, spec, gamma, competing_config)
        stats['misses'] += 1
        cache[key] = result.copy()
        if len(cache) > maxsize:
            cache.popitem(last=False)
        return result

    cached.cache_key = key_for
    cached.cache_info = lambda: dict(**stats, size=len(cache))
    cached.cache_clear = cache.clear
    return cached


def _fit_physical_uncached(train: pd.DataFrame, spec: PhysicalSpec, gamma: float = GAMMA,
                 competing_config: CompetingDamageConfig = DEFAULT_COMPETING_CONFIG,
                 ) -> np.ndarray:
    require_exact_training(train, f"fit_physical[{spec.key}]")
    exact = train.loc[train["is_exact"]].reset_index(drop=True)
    if spec.mode == "linear":
        return _fit_linear_physical(train, spec, gamma)
    if spec.mode != "competing_damage":
        raise ValueError(f"Unknown physical mode: {spec.mode}")

    classic = _fit_linear_physical(train, CLASSIC_WB_SPEC, gamma)
    return water_fit_competing(train,gamma,competing_config,classic,
        _predict_competing_damage,joint_aft_nll,environmental_decrement)


fit_physical = cache_physical_fit(
    _fit_physical_uncached, require_exact_training, lambda: DAMAGE_CAP
)


def predict_physical(df: pd.DataFrame, spec: PhysicalSpec, coef: np.ndarray,
                     gamma: float = GAMMA) -> np.ndarray:
    if spec.mode == "competing_damage":
        return _predict_competing_damage(df, coef, gamma)
    return physical_design(df, spec, gamma) @ coef


def physical_parameter_values(spec: PhysicalSpec, coef: np.ndarray) -> dict[str, float]:

    values = {
        "intercept": np.nan,
        "walker_slope": np.nan,
        "mechanical_log10_damage_scale": np.nan,
        "environment_log10_damage_ratio": np.nan,
        "environment_stress_exponent": np.nan,
        "arrhenius_temperature_sensitivity": np.nan,
        "M_coefficient": np.nan,
        "Ox_coefficient": np.nan,
        "Q_coefficient": np.nan,
    }
    if spec.mode == "competing_damage":
        values.update({
            "intercept": float(coef[0]),
            "mechanical_log10_damage_scale": float(-coef[0]),
            "walker_slope": float(coef[1]),
            "environment_log10_damage_ratio": float(coef[2]),
            "environment_stress_exponent": float(coef[3]),
            "arrhenius_temperature_sensitivity": float(coef[4]),
            "water_log10_multiplier": float(coef[5]) if len(coef)==6 else np.nan,
            "water_multiplier": water_multiplier_value(coef[5]) if len(coef)==6 else np.nan,
        })
    else:
        values["intercept"] = float(coef[0])
        values["walker_slope"] = float(coef[1])
        for channel, value in zip(spec.channels, coef[2:]):
            values[f"{channel}_coefficient"] = float(value)
    return values

