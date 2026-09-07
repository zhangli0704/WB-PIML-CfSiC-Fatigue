"""Compare every result cell without changing either input."""

import argparse
import hashlib
import json
import math
from itertools import zip_longest
from numbers import Real
from pathlib import Path

from openpyxl import load_workbook


WORKBOOK = "PIML_v23_final_results.xlsx"


def compare(reference, candidate, atol=1e-12, rtol=1e-12):
    reference, candidate = Path(reference).resolve(), Path(candidate).resolve()
    directories = reference.is_dir() and candidate.is_dir()
    paths = [p / WORKBOOK if p.is_dir() else p for p in (reference, candidate)]
    report = {
        "reference": str(reference), "candidate": str(candidate),
        "absolute_tolerance": atol, "relative_tolerance": rtol,
        "sheets": [], "compared_cells": 0, "compared_numeric_cells": 0,
        "structure_differences": 0, "type_differences": 0,
        "text_or_other_differences": 0, "numeric_strict_differences": 0,
        "numeric_outside_tolerance": 0, "metadata_differences": 0,
        "max_absolute_difference": 0.0, "max_relative_difference": 0.0,
        "examples": [],
    }

    def difference(kind, location, left, right):
        report[kind] += 1
        if len(report["examples"]) < 15:
            report["examples"].append({
                "kind": kind, "location": location,
                "reference": str(left)[:300], "candidate": str(right)[:300],
            })

    books = []
    try:
        books = [load_workbook(paths[0], read_only=True, data_only=False)]
        books.append(load_workbook(paths[1], read_only=True, data_only=False))
        left, right = books
        if left.sheetnames != right.sheetnames:
            difference("structure_differences", "sheet order", left.sheetnames, right.sheetnames)
        for name in dict.fromkeys(left.sheetnames + right.sheetnames):
            sheets = [book[name] if name in book.sheetnames else None for book in books]
            shapes = [(s.max_row, s.max_column) if s is not None else None for s in sheets]
            item = {"sheet": name, "reference_shape": shapes[0], "candidate_shape": shapes[1]}
            report["sheets"].append(item)
            if shapes[0] != shapes[1]:
                difference("structure_differences", f"{name} dimensions", *shapes)
            if any(s is None for s in sheets):
                continue
            before = report["numeric_strict_differences"]
            for row_number, rows in enumerate(zip_longest(*(s.iter_rows() for s in sheets), fillvalue=()), 1):
                for column, cells in enumerate(zip_longest(*rows), 1):
                    if any(cell is None for cell in cells):
                        continue  # Already reported by the dimension comparison.
                    report["compared_cells"] += 1
                    a, b = (cell.value for cell in cells)
                    types = [(cell.data_type, isinstance(cell.value, bool)) for cell in cells]
                    location = f"{name}!R{row_number}C{column}"
                    if types[0] != types[1]:
                        difference("type_differences", location, types[0], types[1])
                    numeric = all(isinstance(x, Real) and not isinstance(x, bool) for x in (a, b))
                    if numeric:
                        report["compared_numeric_cells"] += 1
                        delta = abs(a - b) if a != b else 0.0
                        relative = delta / max(abs(a), abs(b)) if delta else 0.0
                        for metric, value in (("absolute", delta), ("relative", relative)):
                            key = f"max_{metric}_difference"
                            if value > report[key]:
                                report[key] = value
                                report[f"max_{metric}_location"] = location
                        if a != b:
                            difference("numeric_strict_differences", location, a, b)
                        if not math.isclose(a, b, rel_tol=rtol, abs_tol=atol):
                            report["numeric_outside_tolerance"] += 1
                    elif a != b or type(a) is not type(b):
                        difference("text_or_other_differences", location, a, b)
            item["numeric_strict_differences"] = report["numeric_strict_differences"] - before
    finally:
        for book in books:
            book.close()

    report["metadata_compared"] = directories
    if directories:
        for filename in ("protocol.json", "README_results.md"):
            files = [p / filename for p in (reference, candidate)]
            if not all(p.is_file() for p in files):
                difference("metadata_differences", filename, files[0].is_file(), files[1].is_file())
                continue
            raw = [p.read_bytes() for p in files]
            if filename.endswith(".json"):
                values = [json.dumps(json.loads(x), sort_keys=True, ensure_ascii=False) for x in raw]
                equal = values[0] == values[1]
            else:
                equal = raw[0] == raw[1]
            if not equal:
                difference("metadata_differences", filename,
                           hashlib.sha256(raw[0]).hexdigest(), hashlib.sha256(raw[1]).hexdigest())
    nonnumeric = sum(report[k] for k in (
        "structure_differences", "type_differences", "text_or_other_differences", "metadata_differences"))
    report["strict_equal"] = not (nonnumeric or report["numeric_strict_differences"])
    report["equal_within_tolerance"] = not (nonnumeric or report["numeric_outside_tolerance"])
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("reference", type=Path)
    parser.add_argument("candidate", type=Path)
    parser.add_argument("--json", type=Path, dest="report_path")
    parser.add_argument("--atol", type=float, default=1e-12)
    parser.add_argument("--rtol", type=float, default=1e-12)
    parser.add_argument("--allow-roundoff", action="store_true",
                        help="Accept only numeric differences within tolerance; all other comparisons stay strict.")
    args = parser.parse_args()
    if not all(math.isfinite(x) and x >= 0 for x in (args.atol, args.rtol)):
        parser.error("tolerances must be finite and nonnegative")
    if args.report_path:
        destination = args.report_path.resolve()
        for source in (args.reference.resolve(), args.candidate.resolve()):
            if destination == source or (source.is_dir() and source in destination.parents):
                parser.error("the JSON report must be outside both input directories/files")
    try:
        report = compare(args.reference, args.candidate, args.atol, args.rtol)
    except (OSError, ValueError, KeyError) as exc:
        parser.exit(2, f"Comparison failed: {exc}\n")
    payload = json.dumps(report, ensure_ascii=False, indent=2)
    if args.report_path:
        args.report_path.write_text(payload + "\n", encoding="utf-8")
    print(payload)
    accepted = report["equal_within_tolerance"] if args.allow_roundoff else report["strict_equal"]
    raise SystemExit(0 if accepted else 1)


if __name__ == "__main__":
    main()
